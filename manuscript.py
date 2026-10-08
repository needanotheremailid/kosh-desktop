"""Original shared Markdown subset and inert DOCX/LaTeX writing exports.

Blocks: headings, paragraphs, nested lists, quotes, fenced code, pipe tables,
and equations. Inline: emphasis, strong, code, links, images and limited math.
Raw HTML/TeX is literal. No file/network reads. A caller may resolve an explicit
kosh-asset: identifier to already-authorised image bytes for DOCX figures.
LaTeX export emits source only; compilation/pagination is a separate layer.
"""
from __future__ import annotations

import io
import html
import base64
import binascii
import re
import xml.etree.ElementTree as ET
from urllib.parse import urlsplit


class ManuscriptError(ValueError):
    pass


_LIST = re.compile(r'^( *)(?:([-+*])|(\d+)[.)])\s+(.+)$')
_FENCE = re.compile(r'^\s*(`{3,}|~{3,})(.*)$')
_MANAGED = re.compile(r'^kosh-asset:[A-Za-z0-9_-]{1,100}$')
_MATH_COMMANDS = set(('alpha beta gamma delta epsilon varepsilon zeta eta theta '
                      'vartheta iota kappa lambda mu nu xi pi rho sigma tau upsilon '
                      'phi varphi chi psi omega Gamma Delta Theta Lambda Xi Pi Sigma '
                      'Upsilon Phi Psi Omega frac sqrt sum prod int iint lim sin cos '
                      'tan log ln exp min max cdot times div pm mp le ge leq geq ne neq '
                      'approx equiv in notin subset subseteq supset supseteq cup cap '
                      'infty partial nabla forall exists rightarrow leftarrow '
                      'leftrightarrow Rightarrow Leftarrow Leftrightarrow '
                      'left right big Big hat bar vec dot ddot overline underline '
                      'mathrm mathbf mathit mathcal text operatorname').split())


def safe_math(value):
    """An allowlist grammar, never an arbitrary-TeX opt-in."""
    if not value.strip() or len(value) > 4000 or '^^' in value:
        return False
    depth, index = 0, 0
    while index < len(value):
        char = value[index]
        if char == '\\':
            command = re.match(r'\\([A-Za-z]+|[,;!:{}| ])', value[index:])
            if command is None or (command[1].isalpha() and command[1] not in _MATH_COMMANDS):
                return False
            index += len(command[0])
            continue
        if not (char.isalnum() or char.isspace() or char in '+-*/=()[]{}.,:;|!^_<>'):
            return False
        if char == '{':
            depth += 1
            if depth > 32:
                return False
        elif char == '}':
            depth -= 1
            if depth < 0:
                return False
        index += 1
    return depth == 0


def _safe_link(target):
    if any(ord(char) < 32 for char in target) or any(char in target for char in '\\{}'):
        return False
    try:
        split = urlsplit(target)
    except ValueError:
        return False
    return (split.scheme in {'http', 'https'} and bool(split.netloc)) or (split.scheme == 'mailto' and bool(split.path))


def _closing_paren(text, start):
    depth = 1
    for index in range(start, len(text)):
        if text[index] == '(':
            depth += 1
        elif text[index] == ')':
            depth -= 1
            if depth == 0:
                return index
    return None


def _matching_brackets(text):
    """Map each opening bracket to its own close in one bounded text pass."""
    pending, ends, index = [], {}, 0
    while index < len(text):
        char = text[index]
        if char == '\\':
            index += 2
            continue
        if char == '`':
            marker = re.match(r'`+', text[index:])[0]
            end = text.find(marker, index + len(marker))
            if end >= 0:
                index = end + len(marker)
                continue
        if char == '[':
            pending.append(index)
        elif char == ']' and pending:
            ends[pending.pop()] = index
        index += 1
    return ends


def inline_tokens(text, depth=0):
    if depth > 16:
        return [{'type': 'text', 'text': text}]
    output, plain, index = [], [], 0
    bracket_ends = _matching_brackets(text)
    def flush():
        if plain:
            output.append({'type': 'text', 'text': ''.join(plain)})
            plain.clear()
    while index < len(text):
        char = text[index]
        script_tag = next((tag for tag in ('sup', 'sub') if text.startswith('<' + tag + '>', index)), None)
        if script_tag:
            end = text.find('</' + script_tag + '>', index + 5)
            if end >= 0:
                flush()
                output.append({'type': script_tag, 'children': inline_tokens(text[index + 5:end], depth + 1)})
                index = end + 6
                continue
        if char == '\\' and index + 1 < len(text) and text[index + 1] in '\\`*_{}[]()#+-.!|<>':
            plain.append(text[index + 1])
            index += 2
            continue
        if char == '`':
            marker = re.match(r'`+', text[index:])[0]
            end = text.find(marker, index + len(marker))
            if end >= 0:
                flush()
                output.append({'type': 'code', 'text': text[index + len(marker):end]})
                index = end + len(marker)
                continue
        image = text.startswith('![', index)
        if image or char == '[':
            label_start = index + (2 if image else 1)
            label_end = bracket_ends.get(label_start - 1)
            if label_end is not None and text[label_end + 1:label_end + 2] == '(':
                end = _closing_paren(text, label_end + 2)
                if end is not None:
                    flush()
                    target = text[label_end + 2:end].strip()
                    label = text[label_start:label_end]
                    output.append({'type': 'image' if image else 'link', 'text': label,
                                   'children': inline_tokens(label, depth + 1), 'target': target})
                    index = end + 1
                    continue
        if char == '$' and not text.startswith('$$', index):
            end = text.find('$', index + 1)
            if end >= 0:
                flush()
                expression = text[index + 1:end]
                output.append({'type': 'math', 'text': expression, 'safe': safe_math(expression)})
                index = end + 1
                continue
        marker = next((marker for marker in ('**', '__', '*', '_') if text.startswith(marker, index)), None)
        if marker:
            # Intraword underscores remain ordinary characters (e.g. source_id).
            intraword = marker.startswith('_') and index > 0 and text[index - 1].isalnum()
            end = text.find(marker, index + len(marker)) if not intraword else -1
            if end > index + len(marker):
                flush()
                output.append({'type': 'strong' if len(marker) == 2 else 'emphasis',
                               'children': inline_tokens(text[index + len(marker):end], depth + 1)})
                index = end + len(marker)
                continue
        plain.append(char)
        index += 1
    flush()
    return output


def _cells(line):
    value = line.strip()
    if value.startswith('|'):
        value = value[1:]
    if value.endswith('|') and not value.endswith('\\|'):
        value = value[:-1]
    cells, current, index, code_marker = [], [], 0, ''
    while index < len(value):
        if value[index:index + 2] == '\\|':
            current.append('|')
            index += 2
            continue
        if value[index] == '`':
            marker = re.match(r'`+', value[index:])[0]
            if not code_marker:
                code_marker = marker
            elif marker == code_marker:
                code_marker = ''
            current.append(marker)
            index += len(marker)
            continue
        if value[index] == '|' and not code_marker:
            cells.append(''.join(current).strip())
            current = []
        else:
            current.append(value[index])
        index += 1
    cells.append(''.join(current).strip())
    return cells


def _table_separator(line):
    cells = _cells(line)
    return cells if cells and all(re.fullmatch(r':?-{3,}:?', cell) for cell in cells) else None


def parse_markdown(markdown, _depth=0):
    if not isinstance(markdown, str) or '\x00' in markdown or len(markdown) > 2_000_000:
        raise ManuscriptError('Draft must be text of at most 2,000,000 characters without NUL.')
    if _depth > 16:
        return [{'type': 'paragraph', 'text': markdown}]
    lines, blocks, index = markdown.splitlines(), [], 0
    def special(position):
        line = lines[position]
        return bool(_FENCE.match(line) or _LIST.match(line) or re.match(r'^#{1,6}\s+', line)
                    or re.match(r'^\s*>', line) or line.strip().startswith('$$')
                    or (position + 1 < len(lines) and '|' in line and _table_separator(lines[position + 1])))
    while index < len(lines):
        line = lines[index]
        if not line.strip():
            index += 1
            continue
        fence = _FENCE.match(line)
        if fence:
            content, index = [], index + 1
            while index < len(lines) and not re.fullmatch(r'\s*' + re.escape(fence[1][0]) + '{' + str(len(fence[1])) + r',}\s*', lines[index]):
                content.append(lines[index])
                index += 1
            if index < len(lines):
                index += 1
            blocks.append({'type': 'code', 'text': '\n'.join(content), 'language': fence[2].strip()})
            continue
        display = re.fullmatch(r'\s*\$\$(.+)\$\$\s*', line)
        if display:
            expression = display[1].strip()
            blocks.append({'type': 'math', 'text': expression, 'safe': safe_math(expression)})
            index += 1
            continue
        if line.strip() == '$$':
            content, index = [], index + 1
            while index < len(lines) and lines[index].strip() != '$$':
                content.append(lines[index])
                index += 1
            if index < len(lines):
                index += 1
            expression = '\n'.join(content)
            blocks.append({'type': 'math', 'text': expression, 'safe': safe_math(expression)})
            continue
        heading = re.match(r'^(#{1,6})\s+(.+?)(?:\s+#+)?$', line)
        if heading:
            blocks.append({'type': 'heading', 'level': len(heading[1]), 'text': heading[2]})
            index += 1
            continue
        if re.match(r'^\s*>', line):
            content = []
            while index < len(lines) and re.match(r'^\s*>', lines[index]):
                content.append(re.sub(r'^\s*> ?', '', lines[index]))
                index += 1
            blocks.append({'type': 'quote', 'children': parse_markdown('\n'.join(content), _depth + 1)})
            continue
        marker = _LIST.match(line)
        if marker:
            indent, ordered = len(marker[1]), marker[3] is not None
            items, start = [], int(marker[3]) if ordered else 1
            while index < len(lines):
                match = _LIST.match(lines[index])
                if not match or len(match[1]) != indent or (match[3] is not None) != ordered:
                    break
                item = {'text': match[4], 'children': []}
                index += 1
                children = []
                while index < len(lines):
                    childline = lines[index]
                    if childline.strip() and len(childline) - len(childline.lstrip(' ')) > indent:
                        children.append(childline[indent + 2:])
                        index += 1
                    else:
                        break
                if children:
                    item['children'] = parse_markdown('\n'.join(children), _depth + 1)
                items.append(item)
            blocks.append({'type': 'list', 'ordered': ordered, 'start': start, 'items': items})
            continue
        if index + 1 < len(lines) and '|' in line:
            separator = _table_separator(lines[index + 1])
            header = _cells(line)
            if separator and len(header) == len(separator):
                rows, index = [], index + 2
                while index < len(lines) and lines[index].strip() and '|' in lines[index]:
                    cells = _cells(lines[index])
                    if len(cells) != len(header):
                        break  # Malformed row remains a visible paragraph, never dropped.
                    rows.append(cells)
                    index += 1
                align = ['center' if cell.startswith(':') and cell.endswith(':') else 'right' if cell.endswith(':') else 'left' for cell in separator]
                blocks.append({'type': 'table', 'header': header, 'align': align, 'rows': rows})
                continue
        content = [line]
        index += 1
        while index < len(lines) and lines[index].strip() and not special(index):
            content.append(lines[index])
            index += 1
        blocks.append({'type': 'paragraph', 'text': '\n'.join(content)})
    return blocks


def _literal(value):
    escapes = {'\\': r'\textbackslash{}', '{': r'\{', '}': r'\}', '&': r'\&',
               '%': r'\%', '$': r'\$', '#': r'\#', '_': r'\_',
               '~': r'\textasciitilde{}', '^': r'\textasciicircum{}'}
    return ''.join(escapes.get(char, char) if ord(char) >= 32 else ' ' for char in value)


def _tex_url(target):
    # Percent-encode TeX delimiters; escape percent/hash for macro argument use.
    target = target.replace('{', '%7B').replace('}', '%7D').replace('\\', '%5C').replace('^', '%5E')
    return target.replace('%', r'\%').replace('#', r'\#')


def _latex_inline(tokens):
    parts = []
    for token in tokens:
        kind = token['type']
        if kind in {'sup', 'sub'}:
            parts.append('\\text' + ('super' if kind == 'sup' else 'sub') + 'script{' + _latex_inline(token['children']) + '}')
        elif kind in {'strong', 'emphasis'}:
            command = 'textbf' if kind == 'strong' else 'emph'
            parts.append('\\' + command + '{' + _latex_inline(token['children']) + '}')
        elif kind == 'code':
            parts.append(r'\texttt{' + _literal(token['text']) + '}')
        elif kind == 'math':
            parts.append('$' + token['text'] + '$' if token['safe'] else _literal('$' + token['text'] + '$'))
        elif kind == 'link':
            label = _latex_inline(token['children'])
            parts.append(r'\href{' + _tex_url(token['target']) + '}{' + label + '}' if _safe_link(token['target']) else label + ' (' + _literal(token['target']) + ')')
        elif kind == 'image':
            parts.append(_literal('[Figure: ' + token['text'] + ']'))
        else:
            parts.append(_literal(token['text']))
    return ''.join(parts)


def _latex_blocks(blocks, list_depth=0, ordered_depth=0):
    output = []
    for block in blocks:
        kind = block['type']
        if kind == 'heading':
            command = ('section', 'subsection', 'subsubsection', 'paragraph', 'subparagraph', 'subparagraph')[block['level'] - 1]
            output.append('\\' + command + '*{' + _latex_inline(inline_tokens(block['text'])) + '}')
        elif kind == 'paragraph':
            output.extend([_latex_inline(inline_tokens(block['text'])), ''])
        elif kind == 'list':
            environment = 'enumerate' if block['ordered'] else 'itemize'
            if list_depth >= 4:
                # Standard LaTeX supports four nested list levels. Preserve
                # deeper items visibly without creating a compiler error.
                for item in block['items']:
                    output.append(_literal('- ') + _latex_inline(inline_tokens(item['text'])) + r'\par')
                    output.extend(_latex_blocks(item['children'], list_depth, ordered_depth))
                continue
            output.append(r'\begin{' + environment + '}')
            if block['ordered'] and block['start'] != 1:
                counter = ('enumi', 'enumii', 'enumiii', 'enumiv')[ordered_depth]
                output.append(r'\setcounter{' + counter + '}{' + str(max(0, block['start'] - 1)) + '}')
            for item in block['items']:
                output.append(r'\item ' + _latex_inline(inline_tokens(item['text'])))
                output.extend(_latex_blocks(item['children'], list_depth + 1, ordered_depth + int(block['ordered'])))
            output.append(r'\end{' + environment + '}')
        elif kind == 'quote':
            output.extend([r'\begin{quote}', *_latex_blocks(block['children'], list_depth, ordered_depth), r'\end{quote}'])
        elif kind == 'code':
            # No verbatim environment: an attacker-supplied end marker remains
            # literal even in code. Each source line has a separate paragraph.
            output.append(r'\begin{quote}')
            for line in block['text'].split('\n'):
                output.append(r'\texttt{' + _literal(line) + r'}\par')
            output.append(r'\end{quote}')
        elif kind == 'math':
            output.append(r'\[' + block['text'] + r'\]' if block['safe'] else _literal('$$\n' + block['text'] + '\n$$'))
        elif kind == 'table':
            columns = ''.join({'left': 'l', 'right': 'r', 'center': 'c'}[align] for align in block['align'])
            output.append(r'\begin{tabular}{' + columns + '}')
            output.append(' & '.join(r'\textbf{' + _latex_inline(inline_tokens(cell)) + '}' for cell in block['header']) + r' \\ \hline')
            for row in block['rows']:
                output.append(' & '.join(_latex_inline(inline_tokens(cell)) for cell in row) + r' \\')
            output.append(r'\end{tabular}\par')
    return output


def render_latex(markdown, title=''):
    blocks = parse_markdown(markdown)
    output = [r'\documentclass[11pt]{article}',
              '% Compile with XeLaTeX or LuaLaTeX. Raw document commands are literal.',
              r'\usepackage{fontspec}', r'\usepackage[margin=1in]{geometry}',
              r'\usepackage{amsmath,amssymb}', r'\usepackage{hyperref}',
              r'\setlength{\parindent}{0pt}', r'\setlength{\parskip}{6pt}']
    if title:
        output.extend([r'\title{' + _literal(title) + '}', r'\author{}', r'\date{}'])
    output.append(r'\begin{document}')
    if title:
        output.append(r'\maketitle')
    output.extend(_latex_blocks(blocks))
    output.append(r'\end{document}')
    return '\n'.join(output) + '\n'


def _docx_numbering(document, ordered, start):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    root = document.part.numbering_part.element
    abstract_id = max([int(node.get(qn('w:abstractNumId'))) for node in root.findall(qn('w:abstractNum'))] or [-1]) + 1
    number_id = max([int(node.get(qn('w:numId'))) for node in root.findall(qn('w:num'))] or [0]) + 1
    abstract = OxmlElement('w:abstractNum')
    abstract.set(qn('w:abstractNumId'), str(abstract_id))
    level = OxmlElement('w:lvl')
    level.set(qn('w:ilvl'), '0')
    for tag, value in [('start', '1'), ('numFmt', 'decimal' if ordered else 'bullet'), ('lvlText', '%1.' if ordered else '•')]:
        node = OxmlElement('w:' + tag)
        node.set(qn('w:val'), value)
        level.append(node)
    abstract.append(level)
    first_num = root.find(qn('w:num'))
    if first_num is not None:
        first_num.addprevious(abstract)
    else:
        root.append(abstract)
    num = OxmlElement('w:num')
    num.set(qn('w:numId'), str(number_id))
    reference = OxmlElement('w:abstractNumId')
    reference.set(qn('w:val'), str(abstract_id))
    num.append(reference)
    if ordered:
        override = OxmlElement('w:lvlOverride')
        override.set(qn('w:ilvl'), '0')
        initial = OxmlElement('w:startOverride')
        initial.set(qn('w:val'), str(start))
        override.append(initial)
        num.append(override)
    root.append(num)
    return number_id


def _docx_inline(paragraph, tokens, image_resolver, bold=False, italic=False):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.opc.constants import RELATIONSHIP_TYPE
    from docx.shared import Inches
    for token in tokens:
        kind = token['type']
        if kind in {'sup', 'sub'}:
            before = len(paragraph._p)
            _docx_inline(paragraph, token['children'], image_resolver, bold, italic)
            for node in list(paragraph._p)[before:]:
                for item in ([node] if node.tag == qn('w:r') else node.iter(qn('w:r'))):
                    properties = item.find(qn('w:rPr'))
                    if properties is None:
                        properties = OxmlElement('w:rPr')
                        item.insert(0, properties)
                    alignment = OxmlElement('w:vertAlign')
                    alignment.set(qn('w:val'), 'superscript' if kind == 'sup' else 'subscript')
                    properties.append(alignment)
            continue
        if kind in {'strong', 'emphasis'}:
            _docx_inline(paragraph, token['children'], image_resolver,
                         bold or kind == 'strong', italic or kind == 'emphasis')
            continue
        if kind == 'image':
            data = image_resolver(token['target']) if image_resolver and _MANAGED.fullmatch(token['target']) else None
            if data is None:
                run = paragraph.add_run('[Figure: ' + token['text'] + ']')
            else:
                if not isinstance(data, bytes) or len(data) > 32 * 1024 * 1024:
                    raise ManuscriptError('Managed figure must contain supplied PNG/JPEG/WebP bytes within 32 MiB.')
                if data.startswith(b'RIFF') and data[8:12] == b'WEBP':
                    import pymupdf
                    try:
                        data = pymupdf.Pixmap(data).tobytes('png')
                    except Exception as error:
                        raise ManuscriptError('Managed WebP figure bytes could not be decoded.') from error
                elif not (data.startswith(b'\x89PNG\r\n\x1a\n') or data.startswith(b'\xff\xd8\xff')):
                    raise ManuscriptError('Managed figure must contain supplied PNG/JPEG/WebP bytes within 32 MiB.')
                try:
                    paragraph.add_run().add_picture(io.BytesIO(data), width=Inches(5.5))
                except Exception as error:
                    raise ManuscriptError('Managed figure bytes could not be decoded.') from error
                # Alt text and a visible caption preserve identification.
                picture = paragraph._p.xpath('.//wp:docPr')[-1]
                picture.set('descr', token['text'])
                run = paragraph.add_run('\n' + token['text'])
        elif kind == 'link':
            if _safe_link(token['target']):
                identifier = paragraph.part.relate_to(token['target'], RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
                hyperlink = OxmlElement('w:hyperlink')
                hyperlink.set(qn('r:id'), identifier)
                # Render the label into runs first, then move those real runs
                # into the hyperlink so nested emphasis remains formatted.
                before = len(paragraph._p)
                _docx_inline(paragraph, token['children'], image_resolver, bold, italic)
                for node in list(paragraph._p)[before:]:
                    if node.tag == qn('w:r'):
                        properties = node.find(qn('w:rPr'))
                        if properties is None:
                            properties = OxmlElement('w:rPr')
                            node.insert(0, properties)
                        color = OxmlElement('w:color')
                        color.set(qn('w:val'), '0563C1')
                        properties.append(color)
                        underline = OxmlElement('w:u')
                        underline.set(qn('w:val'), 'single')
                        properties.append(underline)
                    hyperlink.append(node)
                paragraph._p.append(hyperlink)
                continue
            _docx_inline(paragraph, token['children'], image_resolver, bold, italic)
            run = paragraph.add_run(' (' + token['target'] + ')')
        elif kind == 'math':
            equation = _docx_math(token['text']) if token['safe'] else None
            if equation is not None:
                paragraph._p.append(equation)
                continue
            run = paragraph.add_run('$' + token['text'] + '$')
        else:
            value = token['text']
            run = paragraph.add_run(value)
        run.bold, run.italic = bold, italic
        if kind == 'code':
            run.font.name = 'Consolas'
        elif kind == 'math':
            run.font.name = 'Cambria Math'


def _docx_blocks(document, blocks, image_resolver, depth=0, quoted=False):
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt
    for block in blocks:
        kind = block['type']
        if kind == 'heading':
            paragraph = document.add_heading(level=block['level'])
            _docx_inline(paragraph, inline_tokens(block['text']), image_resolver)
        elif kind in {'paragraph', 'math'}:
            paragraph = document.add_paragraph(style='Quote' if quoted else None)
            if kind == 'math':
                equation = _docx_math(block['text'], display=True) if block['safe'] else None
                if equation is not None:
                    paragraph._p.append(equation)
                else:
                    run = paragraph.add_run('$$\n' + block['text'] + '\n$$')
                    run.font.name = 'Cambria Math'
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                _docx_inline(paragraph, inline_tokens(block['text']), image_resolver)
        elif kind == 'list':
            number_id = _docx_numbering(document, block['ordered'], block['start'])
            for item in block['items']:
                paragraph = document.add_paragraph()
                properties = paragraph._p.get_or_add_pPr()
                num_properties = OxmlElement('w:numPr')
                level = OxmlElement('w:ilvl')
                level.set(qn('w:val'), '0')
                number = OxmlElement('w:numId')
                number.set(qn('w:val'), str(number_id))
                num_properties.extend([level, number])
                properties.append(num_properties)
                paragraph.paragraph_format.left_indent = Inches(0.3 * (depth + 1))
                paragraph.paragraph_format.first_line_indent = Inches(-0.15)
                _docx_inline(paragraph, inline_tokens(item['text']), image_resolver)
                _docx_blocks(document, item['children'], image_resolver, depth + 1, quoted)
        elif kind == 'quote':
            _docx_blocks(document, block['children'], image_resolver, depth, True)
        elif kind == 'code':
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.2)
            run = paragraph.add_run(block['text'])
            run.font.name = 'Consolas'
            run.font.size = Pt(9)
        elif kind == 'table':
            table = document.add_table(rows=1, cols=len(block['header']))
            table.style = 'Table Grid'
            for column, value in enumerate(block['header']):
                _docx_inline(table.rows[0].cells[column].paragraphs[0], inline_tokens(value), image_resolver, bold=True)
            for values in block['rows']:
                cells = table.add_row().cells
                for column, value in enumerate(values):
                    _docx_inline(cells[column].paragraphs[0], inline_tokens(value), image_resolver)
            for row in table.rows:
                for column, cell in enumerate(row.cells):
                    cell.paragraphs[0].alignment = {'left': WD_ALIGN_PARAGRAPH.LEFT, 'right': WD_ALIGN_PARAGRAPH.RIGHT, 'center': WD_ALIGN_PARAGRAPH.CENTER}[block['align'][column]]


def render_docx(markdown, title='', image_resolver=None):
    from docx import Document
    blocks = parse_markdown(markdown)
    document = Document()
    document.core_properties.title = title
    document.core_properties.author = ''
    document.core_properties.last_modified_by = ''
    document.core_properties.comments = ''
    if title:
        document.add_heading(title, 0)
    _docx_blocks(document, blocks, image_resolver)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


_MATH_SYMBOLS = dict(zip(
    'alpha beta gamma delta epsilon varepsilon zeta eta theta vartheta iota kappa lambda mu nu xi pi rho sigma tau upsilon phi varphi chi psi omega Gamma Delta Theta Lambda Xi Pi Sigma Upsilon Phi Psi Omega'.split(),
    'α β γ δ ε ϵ ζ η θ ϑ ι κ λ μ ν ξ π ρ σ τ υ φ ϕ χ ψ ω Γ Δ Θ Λ Ξ Π Σ Υ Φ Ψ Ω'.split()))
_MATH_OPERATORS = dict(zip(
    'cdot times div pm mp le leq ge geq ne neq approx equiv in notin subset subseteq supset supseteq cup cap infty partial nabla forall exists rightarrow leftarrow leftrightarrow Rightarrow Leftarrow Leftrightarrow sum prod int iint'.split(),
    '· × ÷ ± ∓ ≤ ≤ ≥ ≥ ≠ ≠ ≈ ≡ ∈ ∉ ⊂ ⊆ ⊃ ⊇ ∪ ∩ ∞ ∂ ∇ ∀ ∃ → ← ↔ ⇒ ⇐ ⇔ ∑ ∏ ∫ ∬'.split()))


class _MathML:
    """Small original formula parser; unsupported syntax falls back to text."""
    def __init__(self, value):
        self.tokens = re.findall(r'\\[A-Za-z]+|\\.|\d+(?:\.\d+)?|[^\W\d_]+|[^\s]', value, re.UNICODE)
        self.pos = 0
        if len(self.tokens) > 1000:
            raise ValueError('Formula exceeds the native layout token bound.')

    def group(self, nested=False, closing='}'):
        values = []
        while self.pos < len(self.tokens) and self.tokens[self.pos] != closing:
            base = self.atom()
            lower, upper = None, None
            while self.pos < len(self.tokens) and self.tokens[self.pos] in {'_', '^'}:
                marker = self.tokens[self.pos]
                self.pos += 1
                script = self.atom()
                if marker == '_':
                    if lower is not None:
                        raise ValueError('Duplicate subscript.')
                    lower = script
                else:
                    if upper is not None:
                        raise ValueError('Duplicate superscript.')
                    upper = script
            if lower is not None and upper is not None:
                base = '<msubsup>' + base + lower + upper + '</msubsup>'
            elif lower is not None:
                base = '<msub>' + base + lower + '</msub>'
            elif upper is not None:
                base = '<msup>' + base + upper + '</msup>'
            values.append(base)
        if nested:
            if self.pos >= len(self.tokens) or self.tokens[self.pos] != closing:
                raise ValueError('Missing formula brace.')
            self.pos += 1
        return '<mrow>' + ''.join(values) + '</mrow>'

    def atom(self):
        if self.pos >= len(self.tokens):
            raise ValueError('Missing formula atom.')
        token = self.tokens[self.pos]
        self.pos += 1
        if token == '{':
            return self.group(True)
        if token in {'}', '_', '^'}:
            raise ValueError('Unsupported formula token order.')
        if token.startswith('\\'):
            command = token[1:]
            if command == 'frac':
                numerator, denominator = self.atom(), self.atom()
                return '<mfrac>' + numerator + denominator + '</mfrac>'
            if command == 'sqrt':
                if self.pos < len(self.tokens) and self.tokens[self.pos] == '[':
                    self.pos += 1
                    degree = self.group(True, ']')
                    return '<mroot>' + self.atom() + degree + '</mroot>'
                return '<msqrt>' + self.atom() + '</msqrt>'
            if command in {'left', 'right', 'big', 'Big'}:
                return self.atom()
            if command in {'mathrm', 'mathbf', 'mathit'}:
                variant = {'mathrm': 'normal', 'mathbf': 'bold', 'mathit': 'italic'}[command]
                return '<mstyle mathvariant="' + variant + '">' + self.atom() + '</mstyle>'
            if command in _MATH_SYMBOLS:
                return '<mi>' + _MATH_SYMBOLS[command] + '</mi>'
            if command in _MATH_OPERATORS:
                return '<mo>' + _MATH_OPERATORS[command] + '</mo>'
            if command in {'sin', 'cos', 'tan', 'log', 'ln', 'exp', 'min', 'max', 'lim'}:
                return '<mi mathvariant="normal">' + command + '</mi>'
            if command in ',;!: ':
                return '<mspace width="0.2em"/>'
            if command in '{}|':
                return '<mo>' + html.escape(command) + '</mo>'
            raise ValueError('Unsupported native formula command.')
        tag = 'mn' if re.fullmatch(r'\d+(?:\.\d+)?', token) else 'mi' if token.isalpha() else 'mo'
        return '<' + tag + '>' + html.escape(token) + '</' + tag + '>'


def _mathml_content(value):
    parser = _MathML(value)
    content = parser.group()
    if parser.pos != len(parser.tokens):
        raise ValueError('Unconsumed formula tokens.')
    return content


def _omml_nodes(node, variant=None):
    """Translate only the elements produced by our bounded formula parser."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    kind = node.tag
    if kind in {'mrow', 'mstyle'}:
        if kind == 'mstyle':
            variant = node.get('mathvariant', variant)
        return [element for child in node for element in _omml_nodes(child, variant)]
    if kind in {'mi', 'mn', 'mo', 'mspace'}:
        run = OxmlElement('m:r')
        math_properties = OxmlElement('m:rPr')
        style = OxmlElement('m:sty')
        style.set(qn('m:val'), {'normal': 'p', 'bold': 'b', 'italic': 'i'}.get(node.get('mathvariant') or variant, 'i' if kind == 'mi' else 'p'))
        math_properties.append(style)
        run.append(math_properties)
        properties = OxmlElement('w:rPr')
        font = OxmlElement('w:rFonts')
        font.set(qn('w:ascii'), 'Cambria Math')
        font.set(qn('w:hAnsi'), 'Cambria Math')
        properties.append(font)
        run.append(properties)
        text = OxmlElement('m:t')
        text.text = ' ' if kind == 'mspace' else node.text or ''
        text.set(qn('xml:space'), 'preserve')
        run.append(text)
        return [run]
    schemas = {'mfrac': ('m:f', ('m:num', 'm:den')), 'msup': ('m:sSup', ('m:e', 'm:sup')),
               'msub': ('m:sSub', ('m:e', 'm:sub')), 'msubsup': ('m:sSubSup', ('m:e', 'm:sub', 'm:sup'))}
    if kind in schemas:
        tag, arguments = schemas[kind]
        element = OxmlElement(tag)
        if kind == 'mfrac':
            properties = OxmlElement('m:fPr')
            fraction_type = OxmlElement('m:type')
            fraction_type.set(qn('m:val'), 'bar')
            properties.append(fraction_type)
            element.append(properties)
        for argument, child in zip(arguments, node):
            content = OxmlElement(argument)
            content.extend(_omml_nodes(child, variant))
            element.append(content)
        return [element]
    if kind in {'msqrt', 'mroot'}:
        radical = OxmlElement('m:rad')
        properties = OxmlElement('m:radPr')
        hidden = OxmlElement('m:degHide')
        hidden.set(qn('m:val'), '1' if kind == 'msqrt' else '0')
        properties.append(hidden)
        radical.append(properties)
        degree = OxmlElement('m:deg')
        content = OxmlElement('m:e')
        if kind == 'mroot':
            content.extend(_omml_nodes(node[0], variant))
            degree.extend(_omml_nodes(node[1], variant))
        else:
            content.extend(element for child in node for element in _omml_nodes(child, variant))
        radical.extend([degree, content])
        return [radical]
    raise ValueError('Formula element is outside the native Word subset.')


def _docx_math(value, display=False):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    if not safe_math(value):
        return None
    try:
        content = ET.fromstring(_mathml_content(value))
        equation = OxmlElement('m:oMath')
        equation.extend(_omml_nodes(content))
        if not display:
            return equation
        paragraph = OxmlElement('m:oMathPara')
        properties = OxmlElement('m:oMathParaPr')
        justification = OxmlElement('m:jc')
        justification.set(qn('m:val'), 'center')
        properties.append(justification)
        paragraph.extend([properties, equation])
        return paragraph
    except (ValueError, ET.ParseError, RecursionError):
        return None


def omml_text(element):
    """Bounded linear notation for our supported OMML; None means omitted.

    No TeX execution or numerical evaluation. Unknown structures/properties
    that can change mathematical meaning must not flatten into a false quote.
    """
    math_ns = '{http://schemas.openxmlformats.org/officeDocument/2006/math}'
    word_ns = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
    budget = [4096]
    def read(node, depth=0):
        budget[0] -= 1
        if budget[0] < 0 or depth > 64:
            raise ValueError('Equation exceeds bounded extraction.')
        kind = node.tag
        if kind == math_ns + 'oMathPara':
            if any(child.tag not in {math_ns+'oMathParaPr', math_ns+'oMath'} for child in node):
                raise ValueError('Unsupported equation paragraph.')
            return ' '.join(read(child, depth+1) for child in node if child.tag == math_ns+'oMath')
        if kind in {math_ns+name for name in ('oMath', 'e', 'num', 'den', 'sub', 'sup', 'deg')}:
            return ''.join(read(child, depth+1) for child in node)
        if kind == math_ns+'r':
            if any(child.tag not in {math_ns+'rPr', word_ns+'rPr', math_ns+'t'} for child in node):
                raise ValueError('Unsupported equation run.')
            value = ''.join(child.text or '' for child in node if child.tag == math_ns+'t')
            if any(ord(char)<32 and char not in '\t\r\n' for char in value):
                raise ValueError('Invalid equation text.')
            style = node.find(math_ns+'rPr/'+math_ns+'sty')
            if style is not None and style.get(math_ns+'val') in {'b', 'bi'}:
                return ('bold(' if style.get(math_ns+'val') == 'b' else 'bolditalic(')+value+')'
            return value
        layouts = {'f': ('num','den'), 'sSub': ('e','sub'), 'sSup': ('e','sup'), 'sSubSup': ('e','sub','sup'), 'rad': ('deg','e')}
        name = kind.removeprefix(math_ns)
        if name not in layouts or kind != math_ns+name:
            raise ValueError('Unsupported native equation structure.')
        argument_names = layouts[name]
        property_name = math_ns+name+'Pr'
        if any(child.tag not in {property_name, *(math_ns+arg for arg in argument_names)} for child in node):
            raise ValueError('Unsupported equation argument.')
        arguments = []
        for argument in argument_names:
            values = [child for child in node if child.tag == math_ns+argument]
            if len(values) != 1:
                raise ValueError('Missing or repeated equation argument.')
            arguments.append(read(values[0], depth+1))
        if name == 'f':
            fraction_type = node.find(math_ns+'fPr/'+math_ns+'type')
            if fraction_type is not None and fraction_type.get(math_ns+'val') not in {'bar','skw','lin'}:
                raise ValueError('Unsupported fraction meaning.')
            return '('+arguments[0]+')/('+arguments[1]+')'
        if name == 'rad':
            degree, radicand = arguments
            return 'sqrt('+radicand+')' if not degree or degree == '2' else 'root('+degree+','+radicand+')'
        base = arguments[0]
        if not re.fullmatch(r'[^\W_]+', base, re.UNICODE):
            base = '('+base+')'
        if name == 'sSub':
            return base+'_{'+arguments[1]+'}'
        if name == 'sSup':
            return base+'^{'+arguments[1]+'}'
        return base+'_{'+arguments[1]+'}^{'+arguments[2]+'}'
    try:
        result = read(element)
        return result if len(result) <= 8000 else None
    except (ValueError, RecursionError):
        return None


def _html_math(value, safe, display=False):
    if safe:
        try:
            content = _mathml_content(value)
            return '<math xmlns="http://www.w3.org/1998/Math/MathML" display="' + ('block' if display else 'inline') + '" aria-label="' + html.escape(value, quote=True) + '">' + content + '</math>'
        except (ValueError, RecursionError):
            pass
    tag = 'div' if display else 'span'
    return '<' + tag + ' class="math' + (' equation' if display else '') + '">' + html.escape(value) + '</' + tag + '>'


def _safe_image_url(target):
    if not isinstance(target, str) or any(ord(char) < 32 for char in target) or '\\' in target:
        return False
    if target.startswith('/') and not target.startswith('//') or target.startswith('blob:'):
        return True
    match = re.fullmatch(r'data:image/(?:png|jpeg|webp);base64,([A-Za-z0-9+/]*={0,2})', target)
    if match and len(match[1]) <= 45_000_000:
        try:
            base64.b64decode(match[1], validate=True)
            return True
        except (ValueError, binascii.Error):
            pass
    return False


def _html_inline(tokens, image_url):
    parts = []
    escape = html.escape
    for token in tokens:
        kind = token['type']
        if kind in {'sup', 'sub'}:
            parts.append('<' + kind + '>' + _html_inline(token['children'], image_url) + '</' + kind + '>')
        elif kind in {'strong', 'emphasis'}:
            tag = 'strong' if kind == 'strong' else 'em'
            parts.append('<' + tag + '>' + _html_inline(token['children'], image_url) + '</' + tag + '>')
        elif kind == 'link':
            label = _html_inline(token['children'], image_url)
            if _safe_link(token['target']):
                parts.append('<a href="' + escape(token['target'], quote=True) + '" rel="noopener noreferrer">' + label + '</a>')
            else:
                parts.append(label + ' (' + escape(token['target']) + ')')
        elif kind == 'image':
            target = image_url(token['target']) if image_url and _MANAGED.fullmatch(token['target']) else None
            # A resolver receives only managed asset IDs. Accept same-origin
            # paths or browser-created blob URLs; never an external-image URL.
            if _safe_image_url(target):
                parts.append('<figure><img src="' + escape(target, quote=True) + '" alt="' + escape(token['text'], quote=True) + '" loading="lazy"><figcaption>' + escape(token['text']) + '</figcaption></figure>')
            else:
                parts.append('<span class="figure-placeholder">' + escape('[Figure: ' + token['text'] + ']') + '</span>')
        elif kind == 'code':
            parts.append('<code>' + escape(token['text']) + '</code>')
        elif kind == 'math':
            parts.append(_html_math(token['text'], token['safe']))
        else:
            parts.append(escape(token['text']).replace('\n', '<br>'))
    return ''.join(parts)


def _html_blocks(blocks, image_url):
    output = []
    for block in blocks:
        kind = block['type']
        if kind in {'paragraph', 'heading'}:
            tag = 'p' if kind == 'paragraph' else 'h' + str(block['level'])
            tokens = inline_tokens(block['text'])
            if kind == 'paragraph' and any(token['type'] == 'image' for token in tokens):
                segment = []
                for token in tokens:
                    if token['type'] == 'image':
                        if segment:
                            output.append('<p>' + _html_inline(segment, image_url) + '</p>')
                            segment = []
                        output.append(_html_inline([token], image_url))
                    else:
                        segment.append(token)
                if segment:
                    output.append('<p>' + _html_inline(segment, image_url) + '</p>')
            else:
                output.append('<' + tag + '>' + _html_inline(tokens, image_url) + '</' + tag + '>')
        elif kind == 'list':
            tag = 'ol' if block['ordered'] else 'ul'
            opening = '<ol start="' + str(block['start']) + '">' if block['ordered'] else '<ul>'
            items = ['<li>' + _html_inline(inline_tokens(item['text']), image_url) + _html_blocks(item['children'], image_url) + '</li>' for item in block['items']]
            output.append(opening + ''.join(items) + '</' + tag + '>')
        elif kind == 'quote':
            output.append('<blockquote>' + _html_blocks(block['children'], image_url) + '</blockquote>')
        elif kind == 'code':
            output.append('<pre><code>' + html.escape(block['text']) + '</code></pre>')
        elif kind == 'math':
            output.append(_html_math(block['text'], block['safe'], display=True))
        elif kind == 'table':
            head = ''.join('<th style="text-align:' + block['align'][column] + '">' + _html_inline(inline_tokens(value), image_url) + '</th>' for column, value in enumerate(block['header']))
            rows = []
            for row in block['rows']:
                rows.append('<tr>' + ''.join('<td style="text-align:' + block['align'][column] + '">' + _html_inline(inline_tokens(value), image_url) + '</td>' for column, value in enumerate(row)) + '</tr>')
            output.append('<table><thead><tr>' + head + '</tr></thead><tbody>' + ''.join(rows) + '</tbody></table>')
    return '\n'.join(output)


def render_html(markdown, image_url=None):
    """Return an escaped HTML fragment using the same writing AST as exports."""
    return _html_blocks(parse_markdown(markdown), image_url)
