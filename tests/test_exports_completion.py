"""Synthetic export composition and actual DOCX/PDF package readback."""
import copy
import base64
import csv
import io
import json
import re
from pathlib import Path
import tempfile
import unittest
import unicodedata
import zipfile

from exports import ExportError, export_workspace


DOC_ID = 'a' * 32
FIG_ID = 'b' * 32
NOTE_ID = 'c' * 32
WORK_ID = 'd' * 32
MARKER = '[[source:' + DOC_ID + ':1]]'


class FakeStore:
    def __init__(self, body=None):
        self.document = {'id': DOC_ID, 'workspace_id': WORK_ID, 'name': 'widgets.txt', 'kind': 'txt', 'pages': 1,
                         'metadata': {'title': 'Widgets', 'authors': 'Doe, Jane', 'author_list': [{'family': 'Doe', 'given': 'Jane'}], 'year': '2024', 'type': 'journal_article', 'journal': 'Widget Review', 'journal_abbreviation': 'Widget Rev', 'volume': '3', 'issue': '2', 'pages': '10-20', 'doi': '10.1000/widget_id'}}
        self.snapshot = {'workspace': {'id': WORK_ID, 'title': 'Synthetic workspace'}, 'documents': [self.document],
                         'notes': [{'id': NOTE_ID, 'title': 'Synthetic manuscript', 'body': body or ('# Findings\n\nWidgets (including blue widgets) ' + MARKER + '\n\n| Name | Count |\n| --- | --- |\n| Blue | 3 |')}],
                         'matrix': [], 'chats': []}
        self.reads = []
        self.bytes = {DOC_ID: b'Synthetic widget source'}

    def _snapshot(self, workspace_id, include_history=False):
        if workspace_id != WORK_ID:
            raise ExportError('Wrong workspace.', 404)
        return copy.deepcopy(self.snapshot)

    def _document(self, document_id, workspace_id=None):
        for item in self.snapshot['documents']:
            if item['id'] == document_id and (workspace_id is None or item['workspace_id'] == workspace_id):
                return copy.deepcopy(item)
        raise ExportError('Source not in this workspace.', 404)

    def _bytes(self, document):
        self.reads.append(document['id'])
        if document['id'] not in self.bytes:
            raise ExportError('Original missing.', 409)
        return self.bytes[document['id']]

    def figure(self):
        import pymupdf
        pixmap = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 20, 10), False)
        pixmap.clear_with(180)
        self.bytes[FIG_ID] = pixmap.tobytes('png')
        self.snapshot['documents'].append({'id': FIG_ID, 'workspace_id': WORK_ID, 'name': 'figure.png', 'kind': 'png', 'pages': 1, 'metadata': {}})
        self.snapshot['notes'][0]['body'] += '\n\n![Synthetic figure](kosh-asset:' + FIG_ID + ')'


def query(format='md', **extra):
    return {'workspace_id': WORK_ID, 'note_id': NOTE_ID, 'format': format, **extra}


class ExportCompletionTests(unittest.TestCase):
    def test_pdf_preserves_literal_entities_and_ordinary_ampersands(self):
        import pymupdf
        from exports import _pdf
        source = 'Widgets & units &amp; &lt;script&gt; &amp;amp; &#169;.'
        with pymupdf.open(stream=_pdf(source, '', {}), filetype='pdf') as document:
            text = unicodedata.normalize('NFKC', document[0].get_text()).strip()
        self.assertEqual(text, source)

    def test_clean_markdown_has_manuscript_not_application_instructions(self):
        store = FakeStore()
        data, mime, name = export_workspace(store, query())
        text = data.decode()
        self.assertTrue(text.startswith('# Synthetic manuscript\n'))
        self.assertIn('## Findings', text)
        self.assertIn('Widgets (including blue widgets) (1)', text)
        self.assertIn('## References', text)
        self.assertNotIn('Local research drafts', text)
        self.assertNotIn('Export review notes', text)
        self.assertNotIn('text unit', text)
        self.assertEqual(mime, 'text/markdown; charset=utf-8')

    def test_optional_audit_contains_warnings_and_separate_locators(self):
        data = export_workspace(FakeStore(), query(audit='1'))[0].decode()
        self.assertIn('## Export review notes', data)
        self.assertIn('text unit 1', data)
        self.assertIn('Evidence locations', data)

    def test_selected_draft_excludes_other_notes_and_missing_uncited_originals(self):
        store = FakeStore()
        other = copy.deepcopy(store.document)
        other['id'] = 'e' * 32
        store.snapshot['documents'].append(other)
        store.snapshot['notes'].append({'id': 'f' * 32, 'title': 'Unselected secret', 'body': 'Other text'})
        text = export_workspace(store, query())[0].decode()
        self.assertNotIn('Unselected secret', text)
        self.assertEqual(store.reads, [DOC_ID])

    def test_cited_missing_original_still_refuses_export(self):
        store = FakeStore()
        store.bytes.clear()
        with self.assertRaises(ExportError):
            export_workspace(store, query())

    def test_workspace_numbers_all_notes_and_matrix_consistently(self):
        store = FakeStore()
        store.snapshot['notes'].append({'id': 'f' * 32, 'title': 'Second draft', 'body': MARKER})
        store.snapshot['matrix'].append({'document_id': DOC_ID, 'question': 'Question', 'design': '', 'findings': MARKER, 'limitations': ''})
        text = export_workspace(store, {'workspace_id': WORK_ID, 'format': 'md'})[0].decode()
        self.assertIn('## Synthetic manuscript', text)
        self.assertIn('### Findings', text)
        self.assertIn('## Second draft', text)
        self.assertIn('## Evidence matrix', text)
        self.assertEqual(text.count('1. Doe J.'), 1)

    def test_apa_pdf_preserves_citation_parentheses_and_tables(self):
        import pymupdf
        data, mime, _ = export_workspace(FakeStore(), query('pdf', citation_style='apa'))
        with pymupdf.open(stream=data, filetype='pdf') as document:
            text = '\n'.join(page.get_text() for page in document)
        self.assertIn('(Doe, 2024)', text)
        self.assertIn('(including blue widgets)', text)
        self.assertIn('Blue', text)
        self.assertIn('Count', text)
        self.assertNotIn('text unit', text)
        self.assertEqual(mime, 'application/pdf')

    def test_docx_real_table_and_formatted_reference(self):
        from docx import Document
        data, _, _ = export_workspace(FakeStore(), query('docx', citation_style='apa'))
        document = Document(io.BytesIO(data))
        self.assertEqual(len(document.tables), 1)
        text = '\n'.join(p.text for p in document.paragraphs)
        self.assertIn('(Doe, 2024)', text)
        self.assertNotIn('*Widget Review', text)
        self.assertTrue(any(run.italic for paragraph in document.paragraphs for run in paragraph.runs))

    def test_docx_figures_are_embedded(self):
        store = FakeStore()
        store.figure()
        data = export_workspace(store, query('docx'))[0]
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            self.assertTrue(any(name.startswith('word/media/') for name in archive.namelist()))
        self.assertIn(FIG_ID, store.reads)

    def test_pdf_figures_are_embedded(self):
        import pymupdf
        store = FakeStore()
        store.figure()
        data = export_workspace(store, query('pdf'))[0]
        with pymupdf.open(stream=data, filetype='pdf') as document:
            self.assertGreater(sum(len(page.get_images()) for page in document), 0)
            self.assertIn('Synthetic figure', unicodedata.normalize('NFKC', '\n'.join(page.get_text() for page in document)))

    def test_pdf_math_typesets_supported_notation_without_mutating_source(self):
        import pymupdf
        body = r'Inline $E=mc^2$, index $x_1$ and fraction $\frac{a}{b}$.' + '\n\n' + r'$$\frac{x_1}{y^2}$$'
        store = FakeStore(body)
        original = copy.deepcopy(store.snapshot)
        data = export_workspace(store, query('pdf'))[0]
        with pymupdf.open(stream=data, filetype='pdf') as document:
            text = '\n'.join(page.get_text() for page in document)
        for literal in ('^', '_', r'\frac'):
            self.assertNotIn(literal, text)
        for glyph in 'Emc2x1aby':
            self.assertIn(glyph, text)
        self.assertEqual(store.snapshot, original)
        html = export_workspace(store, query('html'))[0].decode()
        self.assertIn('<msup>', html)
        self.assertIn('<mfrac>', html)
        latex = export_workspace(store, query('tex'))[0].decode()
        self.assertIn('$E=mc^2$', latex)
        self.assertIn(r'\[\frac{x_1}{y^2}\]', latex)

    def test_pdf_typeset_math_cannot_inject_markup(self):
        import pymupdf
        store = FakeStore('Compare $x < y$ and $x > y$.')
        with pymupdf.open(stream=export_workspace(store, query('pdf'))[0], filetype='pdf') as document:
            text = '\n'.join(page.get_text() for page in document)
        self.assertIn('<', text)
        self.assertIn('>', text)
        self.assertEqual(text.count('x'), 2)
        self.assertEqual(len(re.findall(r'\by\b', text)), 2)

    def test_managed_figure_from_other_workspace_is_rejected(self):
        store = FakeStore()
        store.figure()
        store.snapshot['documents'][-1]['workspace_id'] = 'e' * 32
        with self.assertRaises(ExportError):
            export_workspace(store, query('docx'))

    def test_html_is_self_contained_and_inert(self):
        store = FakeStore('<script>bad</script>\n\n$\\frac{x}{y}$ ' + MARKER)
        store.figure()
        data, mime, _ = export_workspace(store, query('html'))
        text = data.decode()
        self.assertIn('&lt;script&gt;bad&lt;/script&gt;', text)
        self.assertNotIn('<script>', text)
        self.assertIn('data:image/png;base64,', text)
        self.assertIn('<math ', text)
        self.assertEqual(mime, 'text/html; charset=utf-8')

    def test_share_zip_is_portable_without_private_snapshot_or_source_papers(self):
        store = FakeStore()
        store.figure()
        data, mime, _ = export_workspace(store, query('share'))
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            self.assertIn('index.html', archive.namelist())
            self.assertIn('README.txt', archive.namelist())
            self.assertIn('figures/' + FIG_ID + '.png', archive.namelist())
            self.assertNotIn('manifest.json', archive.namelist())
            self.assertEqual([name for name in archive.namelist() if name.endswith('.txt')], ['README.txt'])
            readme = archive.read('README.txt').decode('utf-8')
            self.assertIn('Reading copy, not a restorable workspace backup.', readme)
            self.assertIn('Bibliographic metadata is unverified.', readme)
            self.assertIn('Citation style: Vancouver', readme)
            self.assertRegex(readme, r'Exported at: \d{4}-\d{2}-\d{2}T\d{2}:\d{2}:\d{2}\+00:00')
            self.assertIn('src="figures/' + FIG_ID, archive.read('index.html').decode())
        self.assertEqual(mime, 'application/zip')

    def test_texzip_has_managed_figures_and_safe_caption(self):
        store = FakeStore()
        store.figure()
        store.snapshot['notes'][0]['body'] = store.snapshot['notes'][0]['body'].replace('Synthetic figure', r'Caption \\input{file}')
        data = export_workspace(store, query('texzip'))[0]
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            text = archive.read('main.tex').decode()
            self.assertIn('figures/' + FIG_ID + '.png', archive.namelist())
        self.assertIn('\\includegraphics', text)
        self.assertNotIn('\\input{file}', text)
        self.assertIn('\\textbackslash{}', text)

    def test_tex_source_retains_visible_figure_placeholder(self):
        store = FakeStore()
        store.figure()
        text = export_workspace(store, query('tex'))[0].decode()
        self.assertIn('[Figure: Synthetic figure]', text)
        self.assertIn('caption placeholders', text)

    def test_bib_has_stable_key_raw_doi_underscores_and_group_names(self):
        store = FakeStore()
        store.document['metadata']['author_list'] = [{'literal': 'Widget Group'}]
        text = export_workspace(store, query('bib'))[0].decode()
        self.assertIn('@article{Kosh' + DOC_ID, text)
        self.assertIn('doi = {10.1000/widget_id}', text)
        self.assertNotIn(r'widget\_id', text)
        self.assertIn('author = {{Widget Group}}', text)

    def test_metadata_exports_do_not_read_missing_originals(self):
        store = FakeStore()
        store.bytes.clear()
        for format in ('bib', 'ris', 'csljson'):
            with self.subTest(format=format):
                self.assertTrue(export_workspace(store, query(format))[0])
        self.assertEqual(store.reads, [])

    def test_ris_and_csl_json_metadata(self):
        store = FakeStore()
        self.assertIn('DO  - 10.1000/widget_id', export_workspace(store, query('ris'))[0].decode())
        records = json.loads(export_workspace(store, query('csljson'))[0])
        self.assertEqual(records[0]['id'], 'Kosh' + DOC_ID)
        self.assertEqual(records[0]['DOI'], '10.1000/widget_id')
        self.assertEqual(records[0]['type'], 'article-journal')

    def test_csv_formula_guard_checks_leading_whitespace(self):
        store = FakeStore()
        store.snapshot['matrix'] = [{'document_id': DOC_ID, 'question': '  =1+1', 'design': '+test', 'findings': '@call', 'limitations': '-test'}]
        data = export_workspace(store, {'workspace_id': WORK_ID, 'format': 'csv'})[0]
        rows = list(csv.reader(io.StringIO(data.decode('utf-8-sig'))))
        self.assertEqual(rows[1][-4:], ["'  =1+1", "'+test", "'@call", "'-test"])

    def test_unknown_formats_and_missing_selected_notes_fail(self):
        with self.assertRaises(ExportError):
            export_workspace(FakeStore(), query('unknown'))
        with self.assertRaises(ExportError):
            export_workspace(FakeStore(), query(note_id='f' * 32))

    def test_figure_markup_inside_code_is_not_replaced_by_tex_commands(self):
        store = FakeStore()
        store.figure()
        store.snapshot['notes'][0]['body'] += '\n\n```md\n![Synthetic figure](kosh-asset:' + FIG_ID + ')\n```'
        data = export_workspace(store, query('texzip'))[0]
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            text = archive.read('main.tex').decode()
        self.assertEqual(text.count('\\includegraphics'), 1)
        self.assertIn('kosh-asset:', text)

    def test_html_figure_only_cannot_turn_into_bibliographic_source(self):
        store = FakeStore('Plain paragraph')
        store.figure()
        data = export_workspace(store, query('csljson'))[0]
        self.assertEqual(json.loads(data), [])

    def test_metadata_export_normalizes_doi_prefix_and_preserves_raw_url_underscores(self):
        store = FakeStore()
        store.document['metadata']['doi'] = 'https://doi.org/10.1000/widget_id'
        store.document['metadata']['url'] = 'https://example.test/widget_id'
        bib = export_workspace(store, query('bib'))[0].decode()
        self.assertIn('doi = {10.1000/widget_id}', bib)
        self.assertIn('url = {https://example.test/widget_id}', bib)
        ris = export_workspace(store, query('ris'))[0].decode()
        self.assertIn('DO  - 10.1000/widget_id', ris)
        self.assertEqual(json.loads(export_workspace(store, query('csljson'))[0])[0]['DOI'], '10.1000/widget_id')

    def test_matrix_only_audit_does_not_invent_page_one(self):
        store = FakeStore()
        store.snapshot['notes'] = []
        store.snapshot['matrix'] = [{'document_id': DOC_ID, 'question': 'Document-level summary', 'design': '', 'findings': '', 'limitations': ''}]
        store.document['pages'] = 0
        for style, citation in (('vancouver', '(1)'), ('ieee', '[1]'), ('apa', '(Doe, 2024)')):
            with self.subTest(style=style):
                text = export_workspace(store, {'workspace_id': WORK_ID, 'format': 'md', 'audit': '1', 'citation_style': style})[0].decode()
                self.assertIn('widgets.txt ' + citation, text)
                self.assertIn('Document references', text)
                self.assertIn('no source location', text)
                self.assertNotIn('text unit 1', text)
                self.assertNotIn('PDF p. 1', text)
                self.assertNotIn('Unresolved source', text)

    def test_real_note_page_one_survives_matrix_document_reference(self):
        store = FakeStore()
        store.snapshot['matrix'] = [{'document_id': DOC_ID, 'question': 'Summary only', 'design': '', 'findings': '', 'limitations': ''}]
        text = export_workspace(store, {'workspace_id': WORK_ID, 'format': 'md', 'audit': '1'})[0].decode()
        self.assertEqual(text.count('text unit 1'), 1)
        self.assertIn('no source location', text)
        self.assertEqual(text.count('1. Doe J.'), 1)

    def test_audit_keeps_each_derivative_mapping_when_doi_deduplicates_references(self):
        derivative_id = 'e' * 32
        body = 'Original ' + MARKER + ' and derivative [[source:' + derivative_id + ':1]]'
        store = FakeStore(body)
        derivative = copy.deepcopy(store.document)
        derivative.update(id=derivative_id, kind='pdf', name='derived.pdf', pages=1)
        mapping = 'Unverified local OCR derivative of retained source ' + DOC_ID + '; derived-page to original-file-page mapping: 1->5. Review against original images.'
        derivative['metadata']['provenance'] = mapping
        store.snapshot['documents'].append(derivative)
        store.bytes[derivative_id] = b'Synthetic OCR derivative bytes'
        text = export_workspace(store, query('md', audit='1'))[0].decode()
        self.assertEqual(text.count('1. Doe J.'), 1)
        audit = text.split('### Evidence locations')[1]
        self.assertIn(derivative_id + ': PDF p. 1', audit)
        self.assertIn('Source provenance: ' + mapping, audit)
        self.assertNotIn('Source provenance:', export_workspace(store, query('md'))[0].decode())

    def test_matrix_document_reference_across_all_presentation_formats(self):
        import pymupdf
        from docx import Document
        store = FakeStore()
        store.snapshot['notes'] = []
        store.snapshot['matrix'] = [{'document_id': DOC_ID, 'question': 'Summary only', 'design': '', 'findings': '', 'limitations': ''}]
        store.document['pages'] = 0
        for format in ('md', 'docx', 'pdf', 'tex', 'texzip', 'html', 'share'):
            with self.subTest(format=format):
                data = export_workspace(store, {'workspace_id': WORK_ID, 'format': format, 'citation_style': 'apa'})[0]
                if format == 'docx':
                    text = '\n'.join(p.text for p in Document(io.BytesIO(data)).paragraphs)
                elif format == 'pdf':
                    with pymupdf.open(stream=data, filetype='pdf') as document:
                        text = '\n'.join(page.get_text() for page in document)
                elif format in {'texzip', 'share'}:
                    with zipfile.ZipFile(io.BytesIO(data)) as archive:
                        text = archive.read('main.tex' if format == 'texzip' else 'index.html').decode()
                else:
                    text = data.decode()
                self.assertIn('Doe, 2024', text)
                self.assertNotIn('text unit 1', text)
                self.assertNotIn('PDF p. 1', text)
                self.assertNotIn('Unresolved source', text)

    def test_csv_document_reference_obeys_chosen_style_without_locator(self):
        store = FakeStore()
        store.snapshot['matrix'] = [{'document_id': DOC_ID, 'question': '[[reference:' + DOC_ID + ']]', 'design': '', 'findings': '', 'limitations': ''}]
        data = export_workspace(store, {'workspace_id': WORK_ID, 'format': 'csv', 'citation_style': 'apa'})[0]
        rows = list(csv.reader(io.StringIO(data.decode('utf-8-sig'))))
        self.assertEqual(rows[1][-4], '(Doe, 2024)')

    def test_unselected_figures_are_never_read_or_bibliography_exported(self):
        store = FakeStore()
        store.figure()
        store.snapshot['notes'][0]['body'] = 'No figure or citation used.'
        store.bytes.pop(FIG_ID)
        for format in ('md', 'bib', 'ris', 'csljson'):
            with self.subTest(format=format):
                data = export_workspace(store, query(format))[0]
                self.assertNotIn(FIG_ID.encode(), data)
        self.assertNotIn(FIG_ID, store.reads)

    def test_selected_document_reference_raw_bibliographies_need_no_page_or_original_read(self):
        store = FakeStore('[[reference:' + DOC_ID + ']]')
        store.document['pages'] = 0
        store.bytes.clear()
        self.assertIn(('Kosh' + DOC_ID).encode(), export_workspace(store, query('bib'))[0])
        self.assertIn(('Kosh' + DOC_ID).encode(), export_workspace(store, query('ris'))[0])
        self.assertEqual(json.loads(export_workspace(store, query('csljson'))[0])[0]['id'], 'Kosh' + DOC_ID)
        self.assertEqual(store.reads, [])

    def test_real_managed_asset_import_schema_export_and_tamper_refusal(self):
        import pymupdf
        from backend import AppError, Store
        from completion import Completion
        with tempfile.TemporaryDirectory() as directory:
            store = Store(Path(directory) / 'data')
            try:
                workspace = store.dispatch('POST', '/api/workspaces', {'title': 'Synthetic asset export'})['id']
                pixmap = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 4, 3), False)
                pixmap.clear_with(190)
                receipt = Completion(store).dispatch('POST', '/api/assets/import', {'workspace_id': workspace, 'name': 'synthetic.jpg', 'data': base64.b64encode(pixmap.tobytes('jpeg')).decode('ascii')})
                imported = receipt['results'][0]['document']
                self.assertEqual(imported['kind'], 'jpg')
                self.assertEqual(imported['status'], 'no_text')
                self.assertEqual(imported['pages'], 1)
                note = store._save_note({'workspace_id': workspace, 'title': 'Synthetic figure draft', 'body': '![Managed caption](kosh-asset:' + imported['id'] + ')'})
                for format in ('docx', 'pdf', 'html', 'share', 'texzip'):
                    with self.subTest(format=format):
                        data = export_workspace(store, {'workspace_id': workspace, 'note_id': note['id'], 'format': format})[0]
                        self.assertTrue(data)
                self.assertEqual(json.loads(export_workspace(store, {'workspace_id': workspace, 'note_id': note['id'], 'format': 'csljson'})[0]), [])
                audit = export_workspace(store, {'workspace_id': workspace, 'note_id': note['id'], 'format': 'md', 'audit': '1'})[0].decode()
                self.assertNotIn('text unit 1', audit)
                self.assertNotIn('PDF p. 1', audit)
                document = store._document(imported['id'], workspace)
                (store.root / document['path']).write_bytes(b'Tampered synthetic asset')
                with self.assertRaises(AppError):
                    export_workspace(store, {'workspace_id': workspace, 'note_id': note['id'], 'format': 'docx'})
            finally:
                store.close()

    def test_texzip_bracket_caption_does_not_swallow_other_image_or_prose(self):
        store = FakeStore()
        store.figure()
        store.snapshot['notes'][0]['body'] = '![Remote](https://example.test/image.png) Preserved prose. ![Widget [A]](kosh-asset:' + FIG_ID + ')'
        data = export_workspace(store, query('texzip'))[0]
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            text = archive.read('main.tex').decode()
        self.assertEqual(text.count('\\includegraphics'), 1)
        self.assertIn('\\caption{Widget [A]}', text)
        self.assertIn('Preserved prose.', text)
        self.assertIn('[Figure: Remote]', text)


if __name__ == '__main__':
    unittest.main()
