"""Restore originals across extractor drift without validating stale quotations."""
import base64
import io
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
import zipfile

from backend import AppError, Store


class RestoreDrift(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.store = Store(Path(self.temp.name) / 'data')
        self.addCleanup(self.temp.cleanup)
        self.addCleanup(self.store.close)
        self.workspace = self.store.dispatch('POST', '/api/workspaces', {'title': 'Synthetic extractor recovery'})['id']

    def document(self, kind='docx'):
        if kind == 'docx':
            from docx import Document
            original = Document()
            original.add_paragraph('Thirty synthetic widgets were observed.')
            output = io.BytesIO()
            original.save(output)
            raw = output.getvalue()
        elif kind == 'pdf':
            import pymupdf
            original = pymupdf.open()
            original.new_page().insert_text((40, 40), 'Thirty synthetic widgets were observed.')
            raw = original.tobytes()
            original.close()
        else:
            raw = b'Thirty synthetic widgets were observed.'
        doc = self.store._import({'workspace_id': self.workspace, 'files': [{'name': 'widgets.' + kind, 'data': base64.b64encode(raw).decode()}]})['results'][0]['document']
        chat = self.store.dispatch('POST', '/api/ask', {'workspace_id': self.workspace, 'question': 'widgets', 'model': ''})
        note = self.store._save_note({'workspace_id': self.workspace, 'title': 'Preserved draft', 'body': 'Reviewed widgets [[source:' + doc['id'] + ':1]]'})
        return doc, chat, note, raw

    def backup(self, workspace=None):
        return self.store.file_response('/api/backup', {'workspace_id': workspace or self.workspace})[0]

    def restore(self, archive):
        return self.store.dispatch('POST', '/api/restore', {'data': base64.b64encode(archive).decode()})

    def rewrite(self, archive, change):
        target = io.BytesIO()
        with zipfile.ZipFile(io.BytesIO(archive)) as source, zipfile.ZipFile(target, 'w') as output:
            for name in source.namelist():
                data = source.read(name)
                if name == 'manifest.json':
                    manifest = json.loads(data)
                    change(manifest)
                    data = json.dumps(manifest, ensure_ascii=False).encode()
                output.writestr(name, data)
        return target.getvalue()

    def test_pdf_and_docx_drift_restore_originals_notes_and_quarantine_citation(self):
        for kind in ('pdf', 'docx'):
            with self.subTest(kind=kind):
                # Each format has a distinct workspace so citations remain scoped.
                self.workspace = self.store.dispatch('POST', '/api/workspaces', {'title': 'Synthetic ' + kind})['id']
                doc, chat, note, raw = self.document(kind)
                archive = self.backup()
                with patch('backend.extract', return_value=([(1, 'Different current extracted widgets.')], 'Synthetic changed extractor')):
                    restored = self.restore(archive)
                state = self.store._state(restored['workspace_id'])
                self.assertEqual(restored['unverified_citations'], 1)
                self.assertEqual(self.store._bytes(self.store._document(state['documents'][0]['id'])), raw)
                self.assertIn(state['documents'][0]['id'], state['notes'][0]['body'])
                self.assertEqual(state['chats'][0]['citations'], [])
                unverified = state['chats'][0]['audit']['restored_unverified_citations']
                self.assertEqual(unverified[0]['text'], chat['citations'][0]['text'])
                self.assertEqual(unverified[0]['document_id'], state['documents'][0]['id'])
                self.assertTrue(unverified[0]['restored_unverified'])
                self.assertIn('unverified', state['chats'][0]['warning'].lower())
                second = self.restore(self.backup(restored['workspace_id']))
                second_chat = self.store._state(second['workspace_id'])['chats'][0]
                self.assertEqual(second_chat['citations'], [])
                self.assertEqual(second_chat['audit']['restored_unverified_citations'][0]['text'], chat['citations'][0]['text'])

    def test_extraction_failure_does_not_block_original_and_note_recovery(self):
        doc, chat, note, raw = self.document()
        archive = self.backup()
        with patch('backend.extract', side_effect=AppError('Synthetic new extractor failure')):
            restored = self.restore(archive)
        state = self.store._state(restored['workspace_id'])
        self.assertEqual(state['documents'][0]['status'], 'error')
        self.assertEqual(self.store._bytes(self.store._document(state['documents'][0]['id'])), raw)
        self.assertEqual(len(state['notes']), 1)
        self.assertEqual(restored['unverified_citations'], 1)

    def test_identity_name_page_and_plaintext_quote_forgery_remain_hard_errors(self):
        self.document('txt')
        archive = self.backup()
        changes = [('document_id', 'f' * 32), ('name', 'forged.txt'), ('page', 99), ('text', 'Never in the original synthetic bytes.')]
        before = len(self.store._state()['workspaces'])
        for key, value in changes:
            with self.subTest(key=key):
                bad = self.rewrite(archive, lambda manifest: manifest['chats'][0]['citations'][0].update({key: value}))
                with self.assertRaises(AppError):
                    self.restore(bad)
        self.assertEqual(len(self.store._state()['workspaces']), before)

    def test_pdf_page_forgery_is_rejected_even_when_text_extraction_fails(self):
        self.document('pdf')
        bad = self.rewrite(self.backup(), lambda manifest: manifest['chats'][0]['citations'][0].update(page=2))
        with patch('backend.extract', side_effect=AppError('Synthetic new extraction failure')):
            with self.assertRaises(AppError):
                self.restore(bad)

    def test_pdf_quarantine_retains_historical_bounds_when_reader_is_unavailable(self):
        self.document('pdf')
        archive = self.backup()
        with patch('backend.pdf_lib', None), patch('backend.extract', side_effect=AppError('Synthetic PDF reader unavailable')):
            first = self.restore(archive)
            second = self.restore(self.backup(first['workspace_id']))
        chat = self.store._state(second['workspace_id'])['chats'][0]
        self.assertEqual(chat['citations'], [])
        self.assertEqual(chat['audit']['restored_unverified_citations'][0]['restored_page_limit'], 1)
        self.assertEqual(second['unverified_citations'], 1)

    def test_current_matching_citation_is_kept_without_quarantine(self):
        self.document()
        restored = self.restore(self.backup())
        chat = self.store._state(restored['workspace_id'])['chats'][0]
        self.assertEqual(restored['unverified_citations'], 0)
        self.assertEqual(len(chat['citations']), 1)

    def test_long_unicode_quarantine_survives_backup_and_failing_reextract_again(self):
        from docx import Document
        text = '界' * 4000
        original = Document()
        original.add_paragraph(text)
        output = io.BytesIO()
        original.save(output)
        doc = self.store._import({'workspace_id': self.workspace, 'files': [{'name': 'unicode.docx', 'data': base64.b64encode(output.getvalue()).decode()}]})['results'][0]['document']
        self.store.dispatch('POST', '/api/ask', {'workspace_id': self.workspace, 'question': 'Synthetic widgets', 'model': '', 'selected_source': {'document_id': doc['id'], 'page': 1, 'text': text}})
        archive = self.backup()
        with patch('backend.extract', side_effect=AppError('Synthetic changed extractor')):
            first = self.restore(archive)
            second = self.restore(self.backup(first['workspace_id']))
        chat = self.store._state(second['workspace_id'])['chats'][0]
        self.assertEqual(chat['citations'], [])
        self.assertEqual(chat['audit']['restored_unverified_citations'][0]['text'], text)
        self.assertEqual(second['unverified_citations'], 1)

    def test_generated_provenance_links_remap_but_arbitrary_ids_remain_historical(self):
        source, _, _, _ = self.document('txt')
        linked = self.store._import({'workspace_id': self.workspace, 'files': [{'name': 'linked.txt', 'data': base64.b64encode(b'Other synthetic widgets.').decode()}]})['results'][0]['document']
        provenance = ('Unverified local OCR derivative of retained source ' + source['id'] + '; language eng; derived-page to original-file-page mapping: 1->1. '
                      'User-linked unverified bibliography from retained catalogue source ' + source['id'] + '. '
                      'Historical arbitrary ID ' + source['id'] + ' remains literal.')
        self.store.dispatch('POST', '/api/metadata', {'id': linked['id'], 'expected_metadata_version': linked['metadata_version'], 'metadata': {'provenance': provenance}})
        restored = self.restore(self.backup())
        docs = self.store._state(restored['workspace_id'])['documents']
        current_source = next(doc for doc in docs if doc['name'] == 'widgets.txt')
        current_linked = next(doc for doc in docs if doc['name'] == 'linked.txt')
        text = current_linked['metadata']['provenance']
        self.assertIn('retained source ' + current_source['id'], text)
        self.assertIn('retained catalogue source ' + current_source['id'], text)
        self.assertIn('Historical arbitrary ID ' + source['id'], text)


if __name__ == '__main__':
    unittest.main()
