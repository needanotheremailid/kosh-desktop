"""Synthetic storage acceptance: no user documents or real model calls."""
import base64
import io
import json
import tempfile
import unittest
from unittest.mock import patch
import zipfile
from pathlib import Path

from backend import AppError, Store


def encoded(data):
    return base64.b64encode(data).decode("ascii")


class BackendAcceptance(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.store = Store(Path(self.tmp.name) / "data")
        self.workspace = self.store.dispatch("POST", "/api/workspaces", {"title": "Synthetic research"})

    def tearDown(self):
        self.store.close()
        self.tmp.cleanup()

    def add(self, name="example.txt", data=b"Thirty synthetic widgets were observed. No clinical records.", workspace=None):
        receipt = self.store.dispatch("POST", "/api/import", {"workspace_id": (workspace or self.workspace)["id"], "files": [{"name": name, "data": encoded(data)}]})
        return receipt["results"][0]

    def test_copy_search_provenance_and_workspace_isolation(self):
        document = self.add()["document"]
        found = self.store.dispatch("POST", "/api/search", {"workspace_id": self.workspace["id"], "query": "widgets"})["results"]
        self.assertEqual(found[0]["document_id"], document["id"])
        self.assertEqual(found[0]["page"], 1)
        self.assertIn("Thirty synthetic widgets", found[0]["text"])
        other = self.store.dispatch("POST", "/api/workspaces", {"title": "Other"})
        self.assertEqual(self.store.dispatch("POST", "/api/search", {"workspace_id": other["id"], "query": "widgets"})["results"], [])
        with self.assertRaises(AppError):
            self.store.dispatch("POST", "/api/search", {"workspace_id": other["id"], "query": "widgets", "document_ids": [document["id"]]})
        original, content_type, filename = self.store.file_response("/api/file", {"id": document["id"]})
        self.assertEqual(original, b"Thirty synthetic widgets were observed. No clinical records.")
        self.assertEqual(filename, "example.txt")

    def test_dedup_and_individual_error_receipts(self):
        first = self.add()["document"]
        duplicate = self.add("renamed.txt")
        self.assertEqual(duplicate["status"], "duplicate")
        self.assertEqual(duplicate["document"]["id"], first["id"])
        receipt = self.store.dispatch("POST", "/api/import", {"workspace_id": self.workspace["id"], "files": [{"name": "bad.exe", "data": encoded(b"x")}, {"name": "malformed.pdf", "data": encoded(b"not a pdf")}, {"name": "unicode.md", "data": encoded("Café — naïve evidence".encode())}]})
        self.assertEqual([r["status"] for r in receipt["results"]], ["error", "error", "ready"])
        self.assertEqual(receipt["results"][1]["document"]["status"], "error")

    def test_pdf_actual_page_and_image_only_status(self):
        import pymupdf
        pdf = pymupdf.open()
        pdf.new_page().insert_text((40, 40), "Page one synthetic claim")
        pdf.new_page().insert_text((40, 40), "Page two independent widgets")
        document = self.add("pages.pdf", pdf.tobytes())["document"]
        pdf.close()
        found = self.store.dispatch("POST", "/api/search", {"workspace_id": self.workspace["id"], "query": "independent"})["results"]
        self.assertEqual(found[0]["page"], 2)
        page, content_type, filename = self.store.file_response("/api/page", {"id": document["id"], "page": "2"})
        self.assertEqual(content_type, "image/png")
        self.assertTrue(page.startswith(b"\x89PNG"))
        with self.assertRaises(AppError):
            self.store.file_response("/api/page", {"id": document["id"], "page": "3"})
        blank = pymupdf.open()
        blank.new_page()
        self.assertEqual(self.add("image-only.pdf", blank.tobytes())["status"], "no_text")
        blank.close()

    def test_pdf_import_suggests_title_and_doi_as_unverified_metadata(self):
        import pymupdf
        pdf = pymupdf.open()
        pdf.new_page().insert_text((40, 40), "Invented lantern throughput. https://doi.org/10.1234/abcd.5678. Later text")
        pdf.set_metadata({"title": "Invented Lantern Throughput in Synthetic Workshops"})
        receipt = self.add("lantern.pdf", pdf.tobytes())
        self.assertEqual(receipt["suggested_metadata"], ["doi", "title"])
        metadata = receipt["document"]["metadata"]
        self.assertEqual((metadata["title"], metadata["doi"]), ("Invented Lantern Throughput in Synthetic Workshops", "10.1234/abcd.5678"))
        self.assertIn("verify", metadata["provenance"])
        self.assertEqual(receipt["document"]["metadata_version"], 0)
        junk = pymupdf.open()
        junk.new_page().insert_text((40, 40), "No identifiers here")
        junk.set_metadata({"title": "Microsoft Word - final_manuscript.docx"})
        plain = self.add("plain.pdf", junk.tobytes())
        self.assertNotIn("suggested_metadata", plain)
        self.assertEqual(plain["document"]["metadata"], {})
        saved = self.store.dispatch("POST", "/api/metadata", {"id": receipt["document"]["id"], "expected_metadata_version": 0, "metadata": {**metadata, "year": "2024"}})
        self.assertTrue(saved)
        self.assertEqual(self.store._document(receipt["document"]["id"])["metadata_version"], 1)

    def test_docx_body_table_header_and_coverage_notice(self):
        from docx import Document
        docx = Document()
        docx.add_paragraph("Body synthetic text")
        docx.add_table(rows=1, cols=1).cell(0, 0).text = "Table synthetic evidence"
        docx.sections[0].header.paragraphs[0].text = "Header synthetic context"
        buffer = io.BytesIO()
        docx.save(buffer)
        document = self.add("controlled.docx", buffer.getvalue())["document"]
        result = self.store.dispatch("GET", "/api/document?id=" + document["id"], None)
        text = result["text_pages"][0]["text"]
        self.assertIn("Table synthetic evidence", text)
        self.assertIn("Header synthetic context", text)
        self.assertIn("not", result["extraction_notice"].lower())

    def test_note_conflict_and_restart_persistence(self):
        note = self.store.dispatch("POST", "/api/notes", {"workspace_id": self.workspace["id"], "title": "Draft", "body": "First"})
        updated = self.store.dispatch("POST", "/api/notes", {"workspace_id": self.workspace["id"], "id": note["id"], "version": 1, "title": "Draft", "body": "Second"})
        self.assertEqual(updated["version"], 2)
        with self.assertRaises(AppError) as error:
            self.store.dispatch("POST", "/api/notes", {"workspace_id": self.workspace["id"], "id": note["id"], "version": 1, "title": "Draft", "body": "Stale"})
        self.assertEqual(error.exception.status, 409)
        self.store.close()
        self.store = Store(Path(self.tmp.name) / "data")
        self.assertEqual(self.store.dispatch("GET", "/api/state", None)["notes"][0]["body"], "Second")

    def test_matrix_conflict_and_foreign_document(self):
        document = self.add()["document"]
        fields = {"workspace_id": self.workspace["id"], "document_id": document["id"], "question": "Widget question", "design": "Synthetic", "findings": "Thirty widgets", "limitations": "Fixture only"}
        row = self.store.dispatch("POST", "/api/matrix", fields)
        with self.assertRaises(AppError) as error:
            self.store.dispatch("POST", "/api/matrix", {**fields, "id": row["id"], "version": 0})
        self.assertEqual(error.exception.status, 409)
        other = self.store.dispatch("POST", "/api/workspaces", {"title": "Other"})
        with self.assertRaises(AppError):
            self.store.dispatch("POST", "/api/matrix", {**fields, "workspace_id": other["id"]})

    def test_archive_retains_original_and_notes(self):
        document = self.add()["document"]
        self.store.dispatch("POST", "/api/archive", {"id": document["id"], "archived": True})
        self.assertEqual(self.store.dispatch("POST", "/api/search", {"workspace_id": self.workspace["id"], "query": "widgets"})["results"], [])
        self.assertEqual(self.store.file_response("/api/file", {"id": document["id"]})[0], b"Thirty synthetic widgets were observed. No clinical records.")
        self.store.dispatch("POST", "/api/archive", {"id": document["id"], "archived": False})
        self.assertEqual(len(self.store.dispatch("POST", "/api/search", {"workspace_id": self.workspace["id"], "query": "widgets"})["results"]), 1)

    def test_retrieval_is_labelled_and_unanswerable_has_no_citations(self):
        self.add()
        answer = self.store.dispatch("POST", "/api/ask", {"workspace_id": self.workspace["id"], "question": "How many widgets?", "model": ""})
        self.assertEqual(answer["mode"], "retrieval")
        self.assertIn("excerpt", answer["answer"].lower())
        self.assertIn("Thirty synthetic widgets", answer["citations"][0]["text"])
        missing = self.store.dispatch("POST", "/api/ask", {"workspace_id": self.workspace["id"], "question": "How many nebulae are purple?", "model": ""})
        self.assertEqual(missing["citations"], [])
        self.assertIn("evidence", missing["answer"].lower())

    def test_exports_resolve_citations_and_preserve_entered_metadata(self):
        document = self.add()["document"]
        self.store.dispatch("POST", "/api/metadata", {"id": document["id"], "expected_metadata_version": self.store._document(document["id"])["metadata_version"], "metadata": {"title": "Controlled widgets", "authors": "Example Author", "year": "2024", "doi": "10.example/widget"}})
        self.store.dispatch("POST", "/api/notes", {"workspace_id": self.workspace["id"], "title": "Evidence draft", "body": "Controlled statement [[source:" + document["id"] + ":1]]"})
        for format_ in ("md", "docx", "pdf", "bib", "csv"):
            data, content_type, filename = self.store.file_response("/api/export", {"workspace_id": self.workspace["id"], "format": format_})
            self.assertGreater(len(data), 20)
        md = self.store.file_response("/api/export", {"workspace_id": self.workspace["id"], "format": "md"})[0].decode()
        self.assertNotIn("[[source:", md)
        self.assertIn("Controlled widgets", md)
        self.assertIn("Example Author", md)
        self.assertNotIn("text unit 1", md)
        audit = self.store.file_response('/api/export', {'workspace_id': self.workspace['id'], 'format': 'md', 'audit': '1'})[0].decode()
        self.assertIn('text unit 1', audit)

    def test_backup_restore_fresh_workspace_and_hash_validation(self):
        document = self.add()["document"]
        note = self.store.dispatch("POST", "/api/notes", {"workspace_id": self.workspace["id"], "title": "Saved", "body": "Evidence [[source:" + document["id"] + ":1]]"})
        backup = self.store.file_response("/api/backup", {"workspace_id": self.workspace["id"]})[0]
        restored = self.store.dispatch("POST", "/api/restore", {"data": encoded(backup)})
        self.assertNotEqual(restored["workspace_id"], self.workspace["id"])
        state = self.store.dispatch("GET", "/api/state", None)
        restored_docs = [d for d in state["documents"] if d["workspace_id"] == restored["workspace_id"]]
        restored_notes = [n for n in state["notes"] if n["workspace_id"] == restored["workspace_id"]]
        self.assertNotEqual(restored_docs[0]["id"], document["id"])
        self.assertIn(restored_docs[0]["id"], restored_notes[0]["body"])
        bad = io.BytesIO()
        with zipfile.ZipFile(io.BytesIO(backup)) as original, zipfile.ZipFile(bad, "w") as rewritten:
            for member in original.namelist():
                rewritten.writestr(member, original.read(member) if member == "manifest.json" else b"tampered")
        with self.assertRaises(AppError):
            self.store.dispatch("POST", "/api/restore", {"data": encoded(bad.getvalue())})
        self.assertEqual(len(self.store.dispatch("GET", "/api/state", None)["workspaces"]), 2)

    def test_restore_traversal_rejected_without_workspace_creation(self):
        data = io.BytesIO()
        with zipfile.ZipFile(data, "w") as archive:
            archive.writestr("../outside.txt", "no")
            archive.writestr("manifest.json", json.dumps({"version": 1}))
        with self.assertRaises(AppError):
            self.store.dispatch("POST", "/api/restore", {"data": encoded(data.getvalue())})
        self.assertEqual(len(self.store.dispatch("GET", "/api/state", None)["workspaces"]), 1)

    def test_cloud_inventory_and_unsupported_model_response(self):
        tags = {"models": [{"name": "local:latest", "size": 1024}, {"name": "remote:cloud", "size": 0}, {"name": "remote:120b-cloud", "size": 0}, {"name": "remote:other", "size": 0, "remote_host": "elsewhere"}, {"name": "remote-model:latest", "size": 1, "remote_model": "remote"}]}
        with patch("backend.ollama_request", return_value=tags):
            self.assertEqual([row["name"] for row in self.store.models()["models"]], ["local:latest"])
        self.add()
        with patch.object(self.store, "models", return_value={"models": [{"name": "local:latest", "size": 1024}], "available": True}), patch("backend.ollama_request", return_value={"response": "Unsupported answer [C99]"}):
            answer = self.store.dispatch("POST", "/api/ask", {"workspace_id": self.workspace["id"], "question": "widgets", "model": "local:latest"})
        self.assertEqual(answer["mode"], "retrieval")
        self.assertNotIn("Unsupported answer", answer["answer"])
        self.assertIn("unsupported", answer["warning"].lower())

    def test_local_model_whole_excerpt_budget_and_audit(self):
        for index in range(5):
            self.add("budget-" + str(index) + ".txt", ("widgets " + "界" * 1000 + " " + str(index)).encode())
        captured = {}
        def generate(route, payload, timeout):
            captured.update(payload)
            return {"response": "The source discusses widgets [C1]."}
        with patch.object(self.store, "models", return_value={"models": [{"name": "local:latest", "size": 1024}], "available": True}), patch("backend.ollama_request", side_effect=generate):
            answer = self.store.dispatch("POST", "/api/ask", {"workspace_id": self.workspace["id"], "question": "widgets", "model": "local:latest"})
        self.assertEqual(answer["mode"], "ollama")
        self.assertEqual(captured["options"]["num_ctx"], 8192)
        self.assertLessEqual(len(captured["prompt"].encode()), 6000)
        self.assertIn("界" * 1000, captured["prompt"])
        self.assertEqual(len(answer["audit"]["context"]), 1)
        self.assertEqual(len(answer["audit"]["context"][0]["sha256"]), 64)
        self.assertFalse(answer["audit"]["claim_entailment_verified"])

    def test_csv_formula_cells_and_invalid_reference_warning(self):
        document = self.add()["document"]
        self.store.dispatch("POST", "/api/matrix", {"workspace_id": self.workspace["id"], "document_id": document["id"], "question": "=1+1", "findings": "@formula"})
        csv_bytes = self.store.file_response("/api/export", {"workspace_id": self.workspace["id"], "format": "csv"})[0].decode("utf-8-sig")
        self.assertIn("'=1+1", csv_bytes)
        self.assertIn("'@formula", csv_bytes)
        self.store.dispatch("POST", "/api/notes", {"workspace_id": self.workspace["id"], "title": "Invalid references", "body": "[[source:unresolved:2]] [[source:" + document["id"] + ":99]]"})
        markdown = self.store.file_response("/api/export", {"workspace_id": self.workspace["id"], "format": "md"})[0].decode()
        self.assertNotIn("[[source:", markdown)
        self.assertEqual(markdown.count("[Unresolved source reference:"), 2)

    def test_changed_managed_original_refuses_evidence_and_download(self):
        document = self.add()["document"]
        original = Path(self.tmp.name) / "data" / "originals" / (document["id"] + ".txt")
        original.write_bytes(b"Tampered source")
        with self.assertRaises(AppError):
            self.store.dispatch("POST", "/api/search", {"workspace_id": self.workspace["id"], "query": "widgets"})
        with self.assertRaises(AppError):
            self.store.file_response("/api/file", {"id": document["id"]})

    def test_interrupted_import_reconciles_without_original_deletion(self):
        document = self.add()["document"]
        self.store.db.execute("UPDATE documents SET status='importing' WHERE id=?", (document["id"],))
        self.store.db.execute("DELETE FROM pages WHERE document_id=?", (document["id"],))
        self.store.db.commit()
        self.store.close()
        self.store = Store(Path(self.tmp.name) / "data")
        recovered = self.store.dispatch("GET", "/api/document?id=" + document["id"], None)
        self.assertEqual(recovered["status"], "ready")
        self.assertEqual(recovered["text_pages"][0]["text"], "Thirty synthetic widgets were observed. No clinical records.")

    def test_backup_limit_includes_final_zip_bytes_and_preserves_workspace(self):
        document = self.add()["document"]
        backup = self.store.file_response("/api/backup", {"workspace_id": self.workspace["id"]})[0]
        exact_limit = len(backup)
        with zipfile.ZipFile(io.BytesIO(backup)) as archive:
            content_bytes = sum(member.file_size for member in archive.infolist())
        self.assertLess(content_bytes, exact_limit - 1)
        with patch("backend.BACKUP_LIMIT", exact_limit, create=True):
            published = self.store.file_response("/api/backup", {"workspace_id": self.workspace["id"]})[0]
            self.assertLessEqual(len(published), exact_limit)
        with patch("backend.BACKUP_LIMIT", exact_limit - 1, create=True):
            with self.assertRaises(AppError) as error:
                self.store.file_response("/api/backup", {"workspace_id": self.workspace["id"]})
            self.assertEqual(error.exception.status, 413)
            with self.assertRaises(AppError):
                self.store.dispatch("POST", "/api/restore", {"data": encoded(backup)})
        state = self.store.dispatch("GET", "/api/state", None)
        self.assertEqual(len(state["workspaces"]), 1)
        self.assertEqual(len(state["documents"]), 1)
        self.assertEqual(self.store.file_response("/api/file", {"id": document["id"]})[0], b"Thirty synthetic widgets were observed. No clinical records.")

    def test_backup_collection_bounds_and_escaped_revision_roundtrip(self):
        self.store.dispatch("POST", "/api/notes", {"workspace_id": self.workspace["id"], "title": "First", "body": "Draft"})
        limits = {"documents": 1000, "notes": 1, "matrix": 1000, "chats": 1000, "revisions": 5000}
        with patch("backend.BACKUP_COLLECTION_LIMITS", limits, create=True):
            backup = self.store.file_response("/api/backup", {"workspace_id": self.workspace["id"]})[0]
            self.store.dispatch("POST", "/api/restore", {"data": encoded(backup)})
            self.store.dispatch("POST", "/api/notes", {"workspace_id": self.workspace["id"], "title": "Second", "body": "Other draft"})
            with self.assertRaises(AppError) as error:
                self.store.file_response("/api/backup", {"workspace_id": self.workspace["id"]})
            self.assertEqual(error.exception.status, 413)
        note = self.store.dispatch("POST", "/api/notes", {"workspace_id": self.workspace["id"], "title": "Escaped", "body": "\\" * 200000})
        self.store.dispatch("POST", "/api/notes", {"workspace_id": self.workspace["id"], "id": note["id"], "version": 1, "title": "Escaped", "body": "Updated"})
        backup = self.store.file_response("/api/backup", {"workspace_id": self.workspace["id"]})[0]
        restored = self.store.dispatch("POST", "/api/restore", {"data": encoded(backup)})
        restored_notes = [row for row in self.store.dispatch("GET", "/api/state", None)["notes"] if row["workspace_id"] == restored["workspace_id"]]
        self.assertEqual(len(restored_notes), 3)
        revisions = self.store.db.execute("SELECT payload FROM revisions WHERE entity_id IN (SELECT id FROM notes WHERE workspace_id=?)", (restored["workspace_id"],)).fetchall()
        self.assertEqual(json.loads(revisions[0][0])["body"], "\\" * 200000)

    def test_restore_extracts_each_original_once_for_repeated_citations(self):
        import backend
        self.add()
        for _ in range(3):
            self.store.dispatch("POST", "/api/ask", {"workspace_id": self.workspace["id"], "question": "widgets", "model": ""})
        backup = self.store.file_response("/api/backup", {"workspace_id": self.workspace["id"]})[0]
        with patch("backend.extract", wraps=backend.extract) as extractor:
            restored = self.store.dispatch("POST", "/api/restore", {"data": encoded(backup)})
        self.assertEqual(extractor.call_count, 1)
        self.assertEqual(restored["documents"], 1)

    def test_exact_reimport_repairs_missing_error_original_without_losing_identity(self):
        with patch("backend.atomic_bytes", side_effect=OSError("synthetic disk failure")):
            failed = self.add()
        self.assertEqual(failed["status"], "error")
        document = failed["document"]
        note = self.store.dispatch("POST", "/api/notes", {"workspace_id": self.workspace["id"], "title": "Retained note", "body": "Source [[source:" + document["id"] + ":1]]"})
        for route, query in (("/api/backup", {"workspace_id": self.workspace["id"]}), ("/api/export", {"workspace_id": self.workspace["id"], "format": "md"})):
            with self.assertRaises(AppError) as error:
                self.store.file_response(route, query)
            self.assertIn("reimport", str(error.exception).lower())
        repaired = self.add("same-bytes.txt")
        self.assertEqual(repaired["status"], "ready")
        self.assertEqual(repaired["document"]["id"], document["id"])
        self.assertTrue(repaired["repaired"])
        state = self.store.dispatch("GET", "/api/state", None)
        self.assertEqual(len(state["documents"]), 1)
        self.assertEqual(state["notes"][0]["id"], note["id"])
        self.assertIn(document["id"], state["notes"][0]["body"])
        backup = self.store.file_response("/api/backup", {"workspace_id": self.workspace["id"]})[0]
        self.assertGreater(len(backup), 100)

    def test_exact_reimport_repairs_ready_and_no_text_missing_originals(self):
        for name, content, expected_status in (("ready.txt", b"Synthetic widgets for recovery", "ready"), ("empty.txt", b"", "no_text")):
            document = self.add(name, content)["document"]
            self.assertEqual(document["status"], expected_status)
            managed = Path(self.tmp.name) / "data" / "originals" / (document["id"] + ".txt")
            managed.unlink()
            repaired = self.add(name, content)
            self.assertEqual(repaired["status"], expected_status)
            self.assertTrue(repaired["repaired"])
            self.assertEqual(repaired["document"]["id"], document["id"])
            self.assertEqual(managed.read_bytes(), content)
        found = self.store.dispatch("POST", "/api/search", {"workspace_id": self.workspace["id"], "query": "recovery"})["results"]
        self.assertEqual(len(found), 1)
        self.assertGreater(len(self.store.file_response("/api/backup", {"workspace_id": self.workspace["id"]})[0]), 100)

    def test_current_snapshot_restores_after_history_limit_without_pruning(self):
        note = self.store.dispatch("POST", "/api/notes", {"workspace_id": self.workspace["id"], "title": "Long-lived draft", "body": "Current draft"})
        rows = [("note", note["id"], version, json.dumps({**note, "body": "Historical draft " + str(version), "version": version})) for version in range(1, 5002)]
        self.store.db.executemany("INSERT INTO revisions VALUES(?,?,?,?)", rows)
        self.store.db.execute("UPDATE notes SET version=5002 WHERE id=?", (note["id"],))
        self.store.db.commit()
        with self.assertRaises(AppError) as error:
            self.store.file_response("/api/backup", {"workspace_id": self.workspace["id"]})
        self.assertEqual(error.exception.status, 413)
        snapshot = self.store.file_response("/api/backup", {"workspace_id": self.workspace["id"], "history": "0"})[0]
        with zipfile.ZipFile(io.BytesIO(snapshot)) as archive:
            manifest = json.loads(archive.read("manifest.json"))
        self.assertFalse(manifest["history_included"])
        self.assertEqual(manifest["revisions"], [])
        self.assertIn("retained", manifest["snapshot_note"])
        restored = self.store.dispatch("POST", "/api/restore", {"data": encoded(snapshot)})
        restored_note = next(row for row in self.store.dispatch("GET", "/api/state", None)["notes"] if row["workspace_id"] == restored["workspace_id"])
        self.assertEqual(restored_note["body"], "Current draft")
        self.assertEqual(restored_note["version"], 5002)
        self.assertEqual(self.store.db.execute("SELECT count(*) FROM revisions WHERE entity_id=?", (note["id"],)).fetchone()[0], 5001)
        exported = self.store.file_response("/api/export", {"workspace_id": self.workspace["id"], "format": "md"})[0].decode()
        self.assertIn("Current draft", exported)

    def test_noop_saves_preserve_version_but_stale_versions_still_conflict(self):
        note_fields = {"workspace_id": self.workspace["id"], "title": "Stable draft", "body": "Same content"}
        note = self.store.dispatch("POST", "/api/notes", note_fields)
        unchanged = self.store.dispatch("POST", "/api/notes", {**note_fields, "id": note["id"], "version": 1})
        self.assertEqual(unchanged, note)
        document = self.add()["document"]
        row_fields = {"workspace_id": self.workspace["id"], "document_id": document["id"], "question": "Stable question", "design": "Synthetic", "findings": "Same", "limitations": "Fixture"}
        row = self.store.dispatch("POST", "/api/matrix", row_fields)
        same_row = self.store.dispatch("POST", "/api/matrix", {**row_fields, "id": row["id"], "version": 1})
        self.assertEqual(same_row, row)
        self.assertEqual(self.store.db.execute("SELECT count(*) FROM revisions").fetchone()[0], 0)
        updated_note = self.store.dispatch("POST", "/api/notes", {**note_fields, "id": note["id"], "version": 1, "body": "Changed"})
        updated_row = self.store.dispatch("POST", "/api/matrix", {**row_fields, "id": row["id"], "version": 1, "findings": "Changed"})
        for route, record in (("/api/notes", updated_note), ("/api/matrix", updated_row)):
            with self.assertRaises(AppError) as error:
                self.store.dispatch("POST", route, {**record, "version": 1})
            self.assertEqual(error.exception.status, 409)
        self.assertEqual(self.store.db.execute("SELECT count(*) FROM revisions").fetchone()[0], 2)


if __name__ == "__main__":
    unittest.main()
