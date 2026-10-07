"""User-configured manuscript layout. Never supplies authorship or study claims."""
import io
import re

PROFILES = {'none': 'No manuscript template', 'research': 'Research article', 'review': 'Review article', 'case-report': 'Case report'}
OUTLINES = {
    'research': ['Abstract', 'Introduction', 'Methods', 'Results', 'Discussion', 'Conclusion', 'Declarations'],
    'review': ['Abstract', 'Introduction', 'Methods', 'Findings', 'Discussion', 'Conclusion', 'Declarations'],
    'case-report': ['Abstract', 'Introduction', 'Case presentation', 'Discussion', 'Conclusion', 'Declarations'],
}
DEFAULTS = {'profile': 'none', 'page_size': 'a4', 'font': 'Times New Roman', 'font_size': 12,
            'line_spacing': 2, 'margin_mm': 25.4, 'page_numbers': True, 'line_numbers': False,
            'blinded': False, 'title_page': True, 'authors': '', 'affiliations': '', 'correspondence': '',
            'abstract': '', 'keywords': '', 'running_title': ''}


def validate_template(value):
    if not isinstance(value, dict) or set(value)-set(DEFAULTS):
        raise ValueError('Manuscript template options contain an unsupported field.')
    result = {**DEFAULTS, **value}
    if result['profile'] not in PROFILES or result['page_size'] not in {'a4', 'letter'} or result['font'] not in {'Times New Roman', 'Arial', 'Calibri'}:
        raise ValueError('Choose a listed manuscript profile, page size and font.')
    for field, low, high in [('font_size', 9, 14), ('line_spacing', 1, 3), ('margin_mm', 15, 40)]:
        if type(result[field]) not in {int, float} or not low <= result[field] <= high:
            raise ValueError('Manuscript '+field+' is outside the supported range.')
    for field in ('page_numbers', 'line_numbers', 'blinded', 'title_page'):
        if type(result[field]) is not bool:
            raise ValueError('Manuscript '+field+' must be true or false.')
    for field in ('authors', 'affiliations', 'correspondence', 'abstract', 'keywords', 'running_title'):
        if not isinstance(result[field], str) or len(result[field]) > (12000 if field == 'abstract' else 2000) or any(ord(c)<32 and c not in '\n\t\r' for c in result[field]):
            raise ValueError('Manuscript '+field+' must contain bounded plain text.')
        result[field] = result[field].strip()
    return result


def _plain(text):
    return re.sub(r'([\\`*_{}\[\]()#+.!<>|$])', r'\\\1', text)


def template_frontmatter(options, title):
    if options['profile']=='none': return ''
    values = ['# '+_plain(title), '']
    if not options['blinded']:
        for field in ('authors','affiliations','correspondence'):
            if options[field]: values.extend([_plain(options[field]), ''])
    if options['abstract']: values.extend(['## Abstract', '', _plain(options['abstract']), ''])
    if options['keywords']: values.extend(['**Keywords:** '+_plain(options['keywords']), ''])
    return '\n'.join(values)


def template_markdown(markdown, options, title):
    if options['profile']=='none': return markdown
    # _compose owns the first heading; preserve all user-authored body headings.
    body = markdown.split('\n',1)[1] if markdown.startswith('# ') else markdown
    return template_frontmatter(options,title)+'\n'+body.lstrip()


def template_css(options):
    if options['profile']=='none': return ''
    return 'body {font-family:"'+options['font']+'";font-size:'+str(options['font_size'])+'pt;line-height:'+str(options['line_spacing'])+';} h1,h2,h3 {font-family:"'+options['font']+'";color:#000;} @page {size:'+options['page_size']+';margin:'+str(options['margin_mm'])+'mm;}'


def apply_docx_template(data, options, title):
    if options['profile']=='none': return data
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Mm, Pt, RGBColor
    from docx.enum.text import WD_BREAK
    document = Document(io.BytesIO(data))
    for name in ('Normal','Title', *('Heading '+str(number) for number in range(1,10))):
        style=document.styles[name]
        style.font.name=options['font']
        style.font.color.rgb=RGBColor(0,0,0)
        paragraph_properties=style.element.find(qn('w:pPr'))
        if paragraph_properties is not None:
            for border in list(paragraph_properties.findall(qn('w:pBdr'))):
                paragraph_properties.remove(border)
        properties=style.element.get_or_add_rPr()
        for element in properties:
            for attribute in list(element.attrib):
                if 'theme' in attribute.lower():del element.attrib[attribute]
        style.font.size=Pt(options['font_size'] if name=='Normal' else options['font_size']+2)
        style.paragraph_format.line_spacing=options['line_spacing']
        style.paragraph_format.space_after=Pt(0 if name=='Normal' else 6)
    for section in document.sections:
        section.page_width=Mm(210 if options['page_size']=='a4' else 215.9)
        section.page_height=Mm(297 if options['page_size']=='a4' else 279.4)
        section.top_margin=section.bottom_margin=section.left_margin=section.right_margin=Mm(options['margin_mm'])
        if options['line_numbers']:
            line=OxmlElement('w:lnNumType')
            for key,value in [('countBy','1'),('start','1'),('restart','continuous'),('distance','360')]: line.set(qn('w:'+key),value)
            section._sectPr.append(line)
        if options['page_numbers']:
            footer=section.footer.paragraphs[0]
            footer.alignment=2
            field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),' PAGE ')
            footer._p.append(field)
        if options['running_title']:
            section.header.paragraphs[0].text=options['running_title']
    first=document.paragraphs[0] if document.paragraphs else None
    if first is not None and first.text==title:
        first._p.getparent().remove(first._p)
    anchor=document._element.body[0]
    def add(text,style=None):
        paragraph=document.add_paragraph(text,style)
        anchor.addprevious(paragraph._p)
        return paragraph
    add(title,'Title')
    if not options['blinded']:
        for field in ('authors','affiliations','correspondence'):
            if options[field]: add(options[field])
    if options['title_page']: add('').add_run().add_break(WD_BREAK.PAGE)
    if options['abstract']:
        add('Abstract','Heading 1');add(options['abstract'])
    if options['keywords']: add('Keywords: '+options['keywords'])
    references=False
    for paragraph in document.paragraphs:
        if paragraph.style.name.startswith('Heading'):
            references=paragraph.text.strip().casefold()=='references'
        elif references:
            paragraph.paragraph_format.left_indent=Mm(6.35)
            paragraph.paragraph_format.first_line_indent=Mm(-6.35)
    document.core_properties.author=''
    document.core_properties.last_modified_by=''
    output=io.BytesIO();document.save(output)
    return output.getvalue()


def apply_tex_template(tex, options, title=''):
    if options['profile']=='none': return tex
    from manuscript import _literal, render_latex
    size=options['font_size']
    tex=re.sub(r'\\documentclass\[11pt\]\{article\}',lambda _:r'\documentclass[12pt,'+('a4paper' if options['page_size']=='a4' else 'letterpaper')+']{article}',tex,count=1)
    tex=tex.replace(r'\usepackage[margin=1in]{geometry}',r'\usepackage[margin='+str(options['margin_mm'])+'mm]{geometry}')
    settings=[r'\usepackage{setspace}',r'\setstretch{'+str(options['line_spacing'])+'}']
    # TeX-distributed fonts make this portable without depending on Windows fonts.
    font='texgyretermes' if options['font']=='Times New Roman' else 'texgyreheros'
    settings.append(r'\setmainfont{'+font+'-regular.otf}[BoldFont='+font+'-bold.otf,ItalicFont='+font+'-italic.otf,BoldItalicFont='+font+'-bolditalic.otf]')
    if options['line_numbers']: settings += [r'\usepackage{lineno}',r'\linenumbers']
    if not options['page_numbers']: settings.append(r'\pagestyle{empty}')
    if options['running_title']:
        settings += [r'\usepackage{fancyhdr}',r'\pagestyle{fancy}',r'\fancyhead{}',r'\fancyhead[L]{'+_literal(options['running_title'])+'}']
    tex=tex.replace(r'\begin{document}','\n'.join(settings)+'\n'+r'\begin{document}'+'\n'+r'\fontsize{'+str(size)+'}{'+str(size*1.2)+r'}\selectfont',1)
    if options['title_page'] and title:
        front=render_latex(template_frontmatter({**options,'abstract':'','keywords':''},title))
        front=front.split(r'\begin{document}',1)[1].split(r'\end{document}',1)[0].strip()
        if front not in tex:raise ValueError('The manuscript title page could not be placed; export refused.')
        tex=tex.replace(front,r'\begin{titlepage}'+'\n'+front+'\n'+r'\end{titlepage}',1)
    return tex


def apply_html_template(document, options, title):
    if options['profile']=='none' or not options['title_page']: return document
    from manuscript import render_html
    front=render_html(template_frontmatter({**options,'abstract':'','keywords':''},title))
    return document.replace(front,'<section class="title-page" style="break-after:page;page-break-after:always">'+front+'</section>',1)
